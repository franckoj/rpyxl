"""Model regressions independent of native parsing and serialization."""
from types import SimpleNamespace

import openpyxl
import pytest
from rpyxl.chart import BarChart, Reference
from rpyxl.styles import Color, PatternFill


@pytest.mark.parametrize("sheet", ["Sales", SimpleNamespace(title="Sales")], ids=["title", "worksheet"])
def test_reference_accepts_sheet_title_and_worksheet(sheet):
    reference = Reference(sheet, min_col=2, min_row=1, max_col=3, max_row=4)
    assert reference.sheetname == "Sales"
    assert str(reference) == "'Sales'!$B$1:$C$4"


@pytest.mark.parametrize("from_rows", [False, True])
def test_chart_data_references_match_openpyxl(from_rows):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "O'Brien Sales"
    data = Reference(sheet.title, min_col=1, min_row=1, max_col=3, max_row=4)
    reference_data = openpyxl.chart.Reference(sheet, min_col=1, min_row=1, max_col=3, max_row=4)
    chart = BarChart()
    reference_chart = openpyxl.chart.BarChart()
    chart.add_data(data, titles_from_data=True, from_rows=from_rows)
    reference_chart.add_data(reference_data, titles_from_data=True, from_rows=from_rows)
    assert len(chart.series) == len(reference_chart.series)
    for series, reference_series in zip(chart.series, reference_chart.series):
        assert series.values == reference_series.val.numRef.f
        assert series.title_reference == reference_series.tx.strRef.f
    book.close()


@pytest.mark.parametrize("value", [None, 42, SimpleNamespace(title=None)])
def test_reference_rejects_missing_or_invalid_sheet_title(value):
    with pytest.raises(ValueError, match="worksheet or sheet title"):
        Reference(value)


@pytest.mark.parametrize("rgb", ["FF0000", "FFFF0000", "ff00aa"])
def test_color_and_fill_aliases_match_openpyxl(rgb):
    assert Color(rgb=rgb).rgb == openpyxl.styles.Color(rgb=rgb).rgb.upper()
    fill = PatternFill(fill_type="solid", start_color=rgb)
    reference = openpyxl.styles.PatternFill(fill_type="solid", start_color=rgb)
    assert fill.start_color is fill.fgColor
    assert fill.start_color.rgb == reference.start_color.rgb.upper()
    assert fill.start_color.type == reference.start_color.type
