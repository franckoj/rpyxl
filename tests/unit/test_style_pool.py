"""Style pool value contracts and workbook copy-on-write integration regressions."""

import gc
import weakref
from copy import copy, deepcopy
from dataclasses import fields, is_dataclass
from io import BytesIO
from types import MappingProxyType
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import openpyxl
import pytest
from rpyxl import _metadata, load_workbook
from rpyxl._native_styles import compile_style_records
from rpyxl._style_pool import FrozenStyle, StylePool
from rpyxl.styles import (
    STYLE_INTERNAL_FIELDS,
    Alignment,
    Border,
    Color,
    Font,
    PatternFill,
    Protection,
    Side,
    _style_fields,
)


class CellSlot:
    def __init__(self, pool, value):
        self.pool = pool
        self.value = pool.assign(value)
        self.writable = True
        self.writes = 0

    def check_writable(self):
        if not self.writable:
            raise PermissionError("read-only cell")

    def set_current(self, value):
        self.writes += 1
        self.value = value

    @property
    def view(self):
        return self.pool.view(lambda: self.value, self.set_current, self.check_writable)


@pytest.mark.parametrize(
    "source",
    [
        Font(b=True, sz=14, color="123456"),
        PatternFill(patternType="solid", start_color="ABCDEF", end_color="654321"),
        Border(left=Side(border_style="thin", color=Color(theme=2))),
        Alignment(textRotation=15, wrapText=True, shrinkToFit=True),
        Protection(False, True),
        Color(rgb="#abc123"),
        Side(border_style="dashed", color="123456"),
    ],
)
def test_roundtrip_equality_introspection_and_interning(source):
    pool = StylePool()
    frozen = pool.assign(source)
    assert type(frozen) is FrozenStyle
    assert isinstance(frozen, type(source))
    assert is_dataclass(frozen)
    assert fields(frozen) == fields(source)
    assert frozen == source
    assert source == frozen
    assert frozen.to_model() == source
    assert type(frozen.to_model()) is type(source)
    assert pool.assign(deepcopy(source)) is frozen
    assert frozen.fingerprint is frozen.fingerprint
    assert frozen.fingerprint[0] is type(source)
    assert hash(frozen) == hash(pool.assign(deepcopy(source)))
    assert deepcopy(frozen) is frozen
    assert copy(frozen) is frozen


def test_repeated_assignments_use_cached_freeze_without_reading_scalar_fields(
    monkeypatch,
):
    pool = StylePool()
    source = Font(color=Color(theme=1))
    field_reads = []
    original = Font.__getattribute__
    names = _style_fields(Font)

    def record_reads(self, name):
        if name in names:
            field_reads.append(name)
        return original(self, name)

    monkeypatch.setattr(Font, "__getattribute__", record_reads)
    frozen = pool.assign(source)
    assert field_reads
    field_reads.clear()
    layout_misses = _style_fields.cache_info().misses
    for _ in range(100):
        assert pool.assign(source) is frozen
    assert not field_reads
    assert _style_fields.cache_info().misses == layout_misses
    assert source._style_cache[1] is frozen


@pytest.mark.parametrize(
    "factory", [Font, PatternFill, Border, Alignment, Protection, Color, Side]
)
def test_plain_construction_and_copy_do_not_allocate_tracking(factory, monkeypatch):
    def unexpected_layout(model_type):
        raise AssertionError("Untracked models must not inspect the field layout")

    monkeypatch.setattr("rpyxl.styles._style_fields", unexpected_layout)
    source = factory()
    source.extra = [1]
    del source.extra
    for value in (source, copy(source), deepcopy(source)):
        assert not (vars(value).keys() & STYLE_INTERNAL_FIELDS)


def test_untracked_alias_mutations_and_deletions_remain_ordinary_model_operations():
    source = Font(sz=12, color=Color(theme=1))
    source.b = True
    source.sz = 18
    source.color.tint = 0.25
    source.extra = [1]
    del source.extra
    del source.color
    assert source.bold and source.size == 18
    assert source.color is None
    assert not (vars(source).keys() & STYLE_INTERNAL_FIELDS)
    source.color = Color(theme=2)
    frozen = StylePool().assign(source)
    assert type(frozen) is FrozenStyle
    assert frozen.bold and frozen.size == 18
    assert frozen.color.theme == 2


def test_first_assignment_enables_nested_tracking_before_caching():
    pool = StylePool()
    source = Border(left=Side("thin", Color(theme=1)))
    for model in (source, source.left, source.left.color):
        assert not (vars(model).keys() & STYLE_INTERNAL_FIELDS)
    frozen = pool.assign(source)
    for model in (source, source.left, source.left.color):
        assert vars(model).keys() & STYLE_INTERNAL_FIELDS == STYLE_INTERNAL_FIELDS
        assert model._style_revision == 0
    assert source._style_children == {"left": source.left}
    assert source.left._style_children == {"color": source.left.color}
    source.left.color.tint = 0.5
    changed = pool.assign(source)
    assert changed.left.color.tint == 0.5
    assert frozen.left.color.tint == 0
    replacement = Color(theme=2)
    source.left.color = replacement
    assert not (vars(replacement).keys() & STYLE_INTERNAL_FIELDS)
    replaced = pool.assign(source)
    assert replaced.left.color.theme == 2
    assert vars(replacement).keys() & STYLE_INTERNAL_FIELDS == STYLE_INTERNAL_FIELDS
    replacement.tint = 0.75
    assert pool.assign(source).left.color.tint == 0.75
    assert replaced.left.color.tint == 0


@pytest.mark.parametrize("attribute", ["_xml", "extra", "name"])
def test_opaque_inputs_do_not_activate_tracking_and_can_become_poolable(attribute):
    pool = StylePool()
    source = Font(color=Color(theme=1))
    if attribute == "_xml":
        source._xml = ET.fromstring('<font custom="one"/>')
    else:
        setattr(source, attribute, [1])
    for _ in range(2):
        assigned = pool.assign(source)
        assert type(assigned) is Font
        for model in (source, source.color, assigned, assigned.color):
            assert not (vars(model).keys() & STYLE_INTERNAL_FIELDS)
    if attribute == "name":
        source.name = "Calibri"
    else:
        delattr(source, attribute)
    assert type(pool.assign(source)) is FrozenStyle
    assert source._style_revision == source.color._style_revision == 0


def test_loaded_xml_and_deepcopy_fallback_never_allocate_tracking():
    source = _metadata._read_font(
        ET.fromstring(
            f'<font xmlns="{_metadata.S}"><b/><color theme="2" tint="0.25"/><extension/></font>'
        )
    )
    pool = StylePool()
    first = pool.assign(source)
    source.color.tint = 0.5
    second = pool.assign(source)
    assert first.color.tint == 0.25
    assert second.color.tint == 0.5
    for model in (source, source.color, first, first.color, second, second.color):
        assert not (vars(model).keys() & STYLE_INTERNAL_FIELDS)
    assert source._xml is not first._xml
    assert source.color._xml is not first.color._xml


def test_nested_opaque_models_and_cycles_never_cache_or_activate_the_root():
    pool = StylePool()
    source = Font(color=Color(theme=1))
    source.color.extra = [1]
    assert type(pool.assign(source)) is Font
    assert not (vars(source).keys() & STYLE_INTERNAL_FIELDS)
    assert not (vars(source.color).keys() & STYLE_INTERNAL_FIELDS)
    source.color = source
    for _ in range(2):
        assigned = pool.assign(source)
        assert assigned.color is assigned
        assert not (vars(source).keys() & STYLE_INTERNAL_FIELDS)
        assert not (vars(assigned).keys() & STYLE_INTERNAL_FIELDS)
    source.color = Color(theme=2)
    assert type(pool.assign(source)) is FrozenStyle
    source.color.tint = 0.5
    assert pool.assign(source).color.tint == 0.5
    source.color = source
    assert source._style_cache is None
    assert type(pool.assign(source)) is Font
    assert source._style_cache is None


def test_aliases_and_color_properties():
    pool = StylePool()
    frozen = pool.assign(Font(b=True, i=True, sz=18, u="single", color=Color(theme=3)))
    assert (frozen.b, frozen.i, frozen.sz, frozen.u) == (True, True, 18, "single")
    assert frozen.color.type == "theme"
    assert frozen.color.value == 3
    assert (
        pool.assign(
            Font(
                bold=True,
                italic=True,
                size=18,
                underline="single",
                color=Color(theme=3),
            )
        )
        is frozen
    )
    assert pool.assign(Color(rgb="abc123")).value == "00ABC123"
    assert pool.assign(None) is None
    assert pool.assign(Color()) != pool.assign(Font())


@pytest.mark.parametrize(
    "alias, canonical, value",
    [
        ("b", "bold", True),
        ("i", "italic", True),
        ("sz", "size", 21),
        ("u", "underline", "double"),
    ],
)
def test_source_alias_mutation_increments_revision(alias, canonical, value):
    source = Font()
    pool = StylePool()
    before = pool.assign(source)
    revision = source._style_revision
    setattr(source, alias, value)
    assert source._style_revision > revision
    after = pool.assign(source)
    assert after is not before
    assert getattr(after, canonical) == value
    assert getattr(before, canonical) != value


def test_nested_source_mutations_invalidate_ancestors_without_changing_old_snapshot():
    pool = StylePool()
    source = Border(left=Side("thin", Color(theme=1)))
    first = pool.assign(source)
    revision = source._style_revision
    source.left.color.tint = 0.5
    second = pool.assign(source)
    assert source._style_revision == revision
    assert second is not first
    assert first.left.color.tint == 0
    assert second.left.color.tint == 0.5
    source.left.border_style = "double"
    third = pool.assign(source)
    assert third.left.style == "double"
    assert second.left.style == "thin"
    source.left.color = Color(indexed=5)
    assert pool.assign(source).left.color.indexed == 5


def test_font_and_fill_source_colors_are_independent_snapshots():
    pool = StylePool()
    color = Color(rgb="123456")
    font, fill = Font(color=color), PatternFill(start_color=color)
    first_font, first_fill = pool.assign(font), pool.assign(fill)
    color.rgb = "FFABCDEF"
    assert pool.assign(font).color.rgb == "FFABCDEF"
    assert pool.assign(fill).fgColor.rgb == "FFABCDEF"
    assert first_font.color.rgb == first_fill.fgColor.rgb == "00123456"


def test_cell_nested_edits_update_only_owner_and_keep_nested_views_live():
    pool = StylePool()
    source = Border(left=Side("thin", Color(theme=1)))
    a, b = CellSlot(pool, source), CellSlot(pool, source)
    assert a.value is b.value
    left = a.view.left
    color = left.color
    color.tint = 0.25
    assert a.writes == 1
    assert a.value.left.color.tint == 0.25
    assert b.value.left.color.tint == source.left.color.tint == 0
    a.value = pool.assign(Border(left=Side("double", Color(indexed=2))))
    assert left.style == "double"
    color.tint = 0.5
    assert a.value.left.color.indexed == 2
    assert a.value.left.color.tint == 0.5
    left.border_style = "dashed"
    assert a.value.left.style == "dashed"
    assert b.value.left.style == "thin"


def test_font_color_edits_aliases_and_assignment_of_views():
    pool = StylePool()
    a = CellSlot(pool, Font(color="123456"))
    a.view.b = True
    a.view.sz = 16
    a.view.color.rgb = "FFABCDEF"
    assert a.view == Font(b=True, sz=16, color="FFABCDEF")
    assert isinstance(a.view, Font)
    assert isinstance(a.view.color, Color)
    assert fields(a.view) == fields(Font)
    b = CellSlot(pool, a.view)
    assert b.value is a.value
    a.view.color.tint = 0.5
    assert b.view.color.tint == 0
    b.view.color = a.view.color
    assert b.view.color.tint == 0.5
    a.view.color.tint = 0.75
    assert b.view.color.tint == 0.5
    assert pool.assign(a.view.color) == Color(rgb="FFABCDEF", tint=0.75)


@pytest.mark.parametrize(
    "factory, alias, canonical, value",
    [
        (PatternFill, "patternType", "fill_type", "solid"),
        (Alignment, "textRotation", "text_rotation", 45),
        (Alignment, "wrapText", "wrap_text", True),
        (Alignment, "shrinkToFit", "shrink_to_fit", True),
        (Side, "border_style", "style", "thin"),
    ],
)
def test_non_font_aliases_are_tracked_and_writable_through_views(
    factory, alias, canonical, value
):
    pool = StylePool()
    source = factory()
    before = pool.assign(source)
    setattr(source, alias, value)
    assert getattr(pool.assign(source), canonical) == value
    assert getattr(before, canonical) != value
    cell = CellSlot(pool, factory())
    setattr(cell.view, alias, value)
    assert getattr(cell.value, canonical) == value


@pytest.mark.parametrize("copier", [copy, deepcopy])
def test_copy_of_view_is_an_independent_normal_mutable_tree(copier):
    cell = CellSlot(StylePool(), Border(left=Side("thin", Color(theme=1))))
    result = copier(cell.view)
    assert type(result) is Border
    assert type(result.left) is Side
    assert type(result.left.color) is Color
    result.left.color.tint = 0.5
    result.left.style = "double"
    assert cell.view.left.color.tint == 0
    assert cell.view.left.style == "thin"
    nested = copier(cell.view.left)
    assert type(nested) is Side
    nested.color.theme = 3
    assert cell.view.left.color.theme == 1


def test_nested_fill_aliases_and_unknown_cell_attributes():
    pool = StylePool()
    source = PatternFill("solid", start_color=Color(theme=1))
    a, b = CellSlot(pool, source), CellSlot(pool, source)
    a.view.start_color.tint = 0.25
    assert a.view.fgColor.tint == 0.25
    assert b.view.start_color.tint == 0
    a.view.end_color = a.view.start_color
    assert a.view.bgColor.tint == 0.25
    a.view.custom = ["one"]
    assert type(a.value) is PatternFill
    assert a.view is a.value
    a.view.custom.append("two")
    assert not hasattr(b.view, "custom")
    assert not hasattr(source, "custom")


def test_thawed_models_have_no_mutable_shared_children():
    frozen = StylePool().assign(Font(color=Color(theme=1)))
    a, b = frozen.to_model(), frozen.to_model()
    a.color.tint = 0.5
    assert b.color.tint == frozen.color.tint == 0
    for model in (a, b, a.color, b.color):
        assert not (vars(model).keys() & STYLE_INTERNAL_FIELDS)
    pool = StylePool()
    pool.assign(a)
    pool.assign(b)
    assert a._style_children is not b._style_children
    assert a.color._style_children is not b.color._style_children


def test_shallow_model_copy_does_not_share_tracking_dictionaries():
    pool = StylePool()
    source = Font(color=Color(theme=1))
    frozen = pool.assign(source)
    other = copy(source)
    assert not (vars(other).keys() & STYLE_INTERNAL_FIELDS)
    assert other.color is source.color  # Normal standalone shallow-copy semantics.
    other.color = Color(theme=2)
    assert pool.assign(source) is frozen
    source.color.tint = 0.5
    assert pool.assign(source).color.tint == 0.5
    assert pool.assign(other).color.tint == 0


def test_read_only_check_precedes_any_edit():
    cell = CellSlot(StylePool(), Font(color=Color(theme=1)))
    original = cell.value
    cell.writable = False
    with pytest.raises(PermissionError, match="read-only"):
        cell.view.color.tint = 0.5
    assert cell.value is original
    assert cell.writes == 0


def test_frozen_and_view_dictionaries_cannot_bypass_copy_on_write():
    cell = CellSlot(StylePool(), Font(color=Color(theme=1)))
    frozen = cell.value
    for value in (frozen, frozen.color, cell.view, cell.view.color):
        assert isinstance(vars(value), MappingProxyType)
        with pytest.raises(TypeError):
            vars(value)["color"] = None
        with pytest.raises(AttributeError):
            value._values = ()
    with pytest.raises(AttributeError):
        frozen.bold = True
    with pytest.raises(AttributeError):
        frozen.color.tint = 0.5
    with pytest.raises(AttributeError):
        del frozen.color
    with pytest.raises(AttributeError):
        frozen.fingerprint = ()
    with pytest.raises(AttributeError):
        frozen.__init__(Font, frozen._values)
    assert cell.value is frozen


def test_metadata_serializers_accept_frozen_models():
    pool = StylePool()
    for name, source in [
        ("font", Font(b=True, color=Color(theme=2, tint=0.25))),
        ("fill", PatternFill("solid", fgColor="123456")),
        ("border", Border(left=Side("thin", "ABCDEF"))),
        ("alignment", Alignment(wrapText=True)),
        ("protection", Protection(locked=False)),
    ]:
        frozen = pool.assign(source)
        if name in ("font", "fill", "border"):
            serialize = getattr(_metadata, "_" + name + "_xml")
            actual, expected = serialize(frozen), serialize(source)
        else:
            actual, expected = (
                _metadata._model_xml(name, frozen),
                _metadata._model_xml(name, source),
            )
        assert ET.tostring(actual) == ET.tostring(expected)


def test_original_xml_is_uncached_and_independently_copied_on_every_assignment():
    pool = StylePool()
    source = Font(b=True)
    pool.assign(
        source
    )  # Adding an opaque attribute must invalidate an old clean cache.
    source._xml = ET.fromstring('<font custom="one"><extension/></font>')
    first = pool.assign(source)
    source._xml.set("custom", "two")
    source._xml[0].set("value", "changed")
    second = pool.assign(source)
    assert type(first) is type(second) is Font
    assert first is not source and second is not first
    assert first._xml.get("custom") == "one"
    assert second._xml.get("custom") == "two"
    second._xml.set("custom", "three")
    assert source._xml.get("custom") == "two"
    assert pool.view(lambda: first, lambda value: None, lambda: None) is first
    first.b = False  # The legacy opaque fallback stays writable.
    assert source.b is True


@pytest.mark.parametrize("attribute", ["_xml", "extra"])
def test_nested_opaque_state_invalidates_clean_ancestor_cache(attribute):
    pool = StylePool()
    source = Border(left=Side("thin", Color(theme=1)))
    clean = pool.assign(source)
    setattr(source.left.color, attribute, {"values": [1]})
    first = pool.assign(source)
    getattr(source.left.color, attribute)["values"].append(2)
    second = pool.assign(source)
    assert type(first) is type(second) is Border
    assert getattr(first.left.color, attribute) == {"values": [1]}
    assert getattr(second.left.color, attribute) == {"values": [1, 2]}
    assert not hasattr(clean.left.color, attribute)


@pytest.mark.parametrize("as_extra", [True, False])
def test_mutable_extras_and_field_values_use_uncached_deepcopy(as_extra):
    pool = StylePool()
    source = Font()
    name = "custom" if as_extra else "name"
    setattr(source, name, ["one"])
    first = pool.assign(source)
    getattr(source, name).append("two")
    second = pool.assign(source)
    assert type(first) is type(second) is Font
    assert getattr(first, name) == ["one"]
    assert getattr(second, name) == ["one", "two"]
    assert getattr(first, name) is not getattr(source, name)


def test_unsupported_subclasses_and_cycles_fall_back_safely():
    class CustomFont(Font):
        __slots__ = ("custom",)

    pool = StylePool()
    source = CustomFont(color="123456")
    source.custom = ["one"]
    assigned = pool.assign(source)
    assert type(assigned) is CustomFont and assigned is not source
    source.custom.append("two")
    assert assigned.custom == ["one"]
    source = Font()
    source.color = source
    assigned = pool.assign(source)
    assert type(assigned) is Font
    assert assigned is not source
    assert assigned.color is assigned


def test_repeated_style_access_does_not_construct_discarded_metadata(monkeypatch):
    import rpyxl.workbook as workbook_module
    from rpyxl import Workbook, _compat

    sheet_factory = workbook_module.MetadataSheet
    cell_factory = _compat.CellMetadata
    counts = {"sheet": 0, "cell": 0}

    def make_sheet(*args, **kwargs):
        counts["sheet"] += 1
        return sheet_factory(*args, **kwargs)

    def make_cell(*args, **kwargs):
        counts["cell"] += 1
        return cell_factory(*args, **kwargs)

    monkeypatch.setattr(workbook_module, "MetadataSheet", make_sheet)
    monkeypatch.setattr(_compat, "CellMetadata", make_cell)
    book = Workbook()
    try:
        cell = book.active["A1"]
        source = Font(bold=True)
        for _ in range(20):
            cell.font = source
            assert cell.font.bold
        assert counts == {"sheet": 1, "cell": 1}
    finally:
        book.close()


def test_workbook_default_styles_are_built_once_and_remain_isolated(monkeypatch):
    from rpyxl import Workbook, _compat

    original = _compat.STYLE_DEFAULTS["font"]
    calls = []

    def create_font():
        calls.append(True)
        return original()

    monkeypatch.setitem(_compat.STYLE_DEFAULTS, "font", create_font)
    book = Workbook()
    try:
        sheet = book.active
        for row in range(1, 101):
            assert not sheet.cell(row, 1).font.bold
        assert len(calls) == 1
        sheet["A1"].font.bold = True
        assert sheet["A1"].font.bold
        assert not sheet["A2"].font.bold
        assert not book._style_defaults["font"].bold
        assert len(calls) == 1
    finally:
        book.close()


def test_conditional_formatting_accepts_a_cell_style_view():
    from io import BytesIO

    import openpyxl
    from rpyxl import Workbook
    from rpyxl.formatting.rule import FormulaRule

    book = Workbook()
    book.active["A1"].font = Font(bold=True)
    book.active.conditional_formatting.add(
        "B1", FormulaRule(formula=["TRUE"], font=book.active["A1"].font)
    )
    output = BytesIO()
    try:
        book.save(output)
        loaded = openpyxl.load_workbook(output)
        try:
            formatting = loaded.active.conditional_formatting
            conditions = list(formatting)
            assert len(conditions) == 1
            assert formatting[conditions[0]][0].dxf.font.bold
        finally:
            loaded.close()
    finally:
        book.close()


def test_shallow_copy_of_opaque_style_subclass_preserves_slots():
    class CustomFont(Font):
        __slots__ = ("__private", "tag")

    source = CustomFont()
    source.tag = ["important"]
    source._CustomFont__private = "retained"
    assigned = StylePool().assign(source)
    copied = copy(assigned)
    assert copied.tag is assigned.tag
    assert copied.tag is not source.tag
    assert copied._CustomFont__private == "retained"


def test_tracking_fields_are_explicit_and_never_dataclass_fields():
    source = Font(color=Color(theme=1))
    StylePool().assign(source)
    assert set(vars(source)) - set(_style_fields(Font)) == STYLE_INTERNAL_FIELDS
    assert STYLE_INTERNAL_FIELDS == {
        "_style_revision",
        "_style_children",
        "_style_opaque",
        "_style_cache",
    }
    assert not (set(_style_fields(Font)) & STYLE_INTERNAL_FIELDS)
    assert not (
        set(vars(deepcopy(source))) - set(_style_fields(Font)) - STYLE_INTERNAL_FIELDS
    )


def test_pool_does_not_retain_sources_or_obsolete_frozen_values():
    pool = StylePool()
    source = Font(color="123456")
    frozen = pool.assign(source)
    source_ref, frozen_ref = weakref.ref(source), weakref.ref(frozen)
    del source
    gc.collect()
    assert source_ref() is None
    assert frozen_ref() is frozen
    del frozen
    gc.collect()
    assert frozen_ref() is None
    assert not pool._styles


@pytest.mark.parametrize(
    "field, values",
    [("sz", (12, 12.0)), ("b", (True, 1, 1.0))],
)
def test_interning_preserves_primitive_types_and_consistent_frozen_equality(
    field, values
):
    pool = StylePool()
    frozen = [pool.assign(Font(**{field: value})) for value in values]
    assert [type(getattr(style, field)) for style in frozen] == [
        type(value) for value in values
    ]
    assert [type(getattr(style.to_model(), field)) for style in frozen] == [
        type(value) for value in values
    ]
    assert len({style.fingerprint for style in frozen}) == len(values)
    assert len(set(frozen)) == len(values)
    for value, style in zip(values, frozen):
        equal = pool.assign(Font(**{field: value}))
        assert equal is style
        assert equal == style
        assert hash(equal) == hash(style)
        assert style == Font(**{field: value})


def test_nested_fingerprints_preserve_numeric_types_and_model_alias_equality():
    pool = StylePool()
    integer = pool.assign(Font(color=Color(tint=0)))
    floating = pool.assign(Font(color=Color(tint=0.0)))
    assert integer is not floating
    assert integer != floating
    assert integer.fingerprint != floating.fingerprint
    assert type(integer.color.tint) is int
    assert type(floating.color.tint) is float
    assert type(integer.to_model().color.tint) is int
    assert type(floating.to_model().color.tint) is float
    assert pool.assign(Font(sz=12)) is pool.assign(Font(size=12))
    assert Font(sz=12) == Font(size=12.0)  # Standalone dataclass semantics stay intact.


def test_workbook_numeric_font_alias_assignments_preserve_value_types(rpyxl_wb):
    sheet = rpyxl_wb.active
    sheet["A1"].font = Font(sz=12)
    sheet["B1"].font = Font(size=12.0)
    assert type(sheet["A1"].font.size) is int
    assert type(sheet["B1"].font.sz) is float
    assert sheet._metadata.cells["A1"].font is not sheet._metadata.cells["B1"].font
    assert type(copy(sheet["A1"].font).size) is int
    assert type(deepcopy(sheet["B1"].font).sz) is float


def test_workbook_repeated_font_assignment_shares_frozen_values_and_native_record(
    rpyxl_wb,
):
    sheet = rpyxl_wb.active
    source = Font(b=True, color=Color(theme=2, tint=0.25))
    for row in range(1, 101):
        sheet.cell(row=row, column=1, value=row).font = source
    stored = [sheet._metadata.cells[f"A{row}"].font for row in range(1, 101)]
    assert type(stored[0]) is FrozenStyle
    assert all(font is stored[0] for font in stored)
    assert source._style_cache[1] is stored[0]
    records, assignments = compile_style_records(rpyxl_wb._book)
    assert len(records) == 1
    assert assignments == [(sheet.title, row, 1, 0) for row in range(1, 101)]


def test_workbook_source_mutation_only_affects_subsequent_assignments(rpyxl_wb):
    sheet = rpyxl_wb.active
    source = Font(color=Color(theme=2))
    sheet["A1"].font = sheet["B1"].font = source
    original = sheet._metadata.cells["A1"].font
    source.b = True
    source.color.tint = 0.5
    sheet["C1"].font = source
    assert sheet._metadata.cells["A1"].font is original
    assert sheet._metadata.cells["B1"].font is original
    for coordinate in ("A1", "B1"):
        assert not sheet[coordinate].font.b
        assert sheet[coordinate].font.color.tint == 0
    assert sheet["C1"].font.b
    assert sheet["C1"].font.color.tint == 0.5
    source.color.tint = 0.75  # Child-only mutation must invalidate the source cache.
    sheet["D1"].font = source
    assert sheet["C1"].font.color.tint == 0.5
    assert sheet["D1"].font.color.tint == 0.75


@pytest.mark.parametrize("copier", [copy, deepcopy])
def test_workbook_copy_font_is_independent_normal_mutable_model(rpyxl_wb, copier):
    sheet = rpyxl_wb.active
    sheet["A1"].font = Font(b=True, sz=14, color=Color(theme=2))
    result = copier(sheet["A1"].font)
    assert type(result) is Font
    assert type(result.color) is Color
    result.b = False
    result.sz = 18
    result.color.tint = 0.5
    assert sheet["A1"].font.b
    assert sheet["A1"].font.sz == 14
    assert sheet["A1"].font.color.tint == 0
    sheet["B1"].font = result
    result.color.tint = 0.75
    assert sheet["B1"].font.color.tint == 0.5


def test_workbook_nested_views_and_aliases_only_edit_the_owning_cell(rpyxl_wb):
    sheet = rpyxl_wb.active
    font = Font(color="123456")
    border = Border(left=Side("thin", Color(theme=2)))
    for coordinate in ("A1", "B1"):
        sheet[coordinate].font = font
        sheet[coordinate].border = border
    sheet["C1"].font = sheet["A1"].font
    assert sheet._metadata.cells["C1"].font is sheet._metadata.cells["A1"].font
    sheet["A1"].font.color.rgb = "FFABCDEF"
    sheet["A1"].font.b = True
    sheet["A1"].font.sz = 18
    sheet["A1"].border.left.color.tint = 0.5
    sheet["A1"].border.left.border_style = "double"
    assert sheet["A1"].font.bold
    assert sheet["A1"].font.size == 18
    assert sheet["A1"].font.color.rgb == "FFABCDEF"
    assert sheet["A1"].border.left.style == "double"
    assert sheet["A1"].border.left.color.tint == 0.5
    for coordinate in ("B1", "C1"):
        assert not sheet[coordinate].font.bold
        assert sheet[coordinate].font.size == 11.0
        assert sheet[coordinate].font.color.rgb == "00123456"
    assert sheet["B1"].border.left.style == "thin"
    assert sheet["B1"].border.left.color.tint == 0
    assert font.color.rgb == "00123456"
    assert border.left.color.tint == 0


def test_workbook_cow_save_preserves_components_values_and_number_formats(rpyxl_wb):
    sheet = rpyxl_wb.active
    font = Font(b=True, color="123456")
    fill = PatternFill("solid", start_color="654321")
    border = Border(left=Side("thin", Color(theme=2, tint=0.25)))
    alignment = Alignment(wrapText=False, horizontal="center")
    protection = Protection(locked=True)
    for coordinate, value in (("A1", 12.5), ("B1", 37.25)):
        cell = sheet[coordinate]
        cell.value = value
        cell.number_format = '0.0000 "units"'
        cell.font, cell.fill, cell.border = font, fill, border
        cell.alignment, cell.protection = alignment, protection
    sheet["A1"].font.color.rgb = "FFABCDEF"
    sheet["A1"].font.sz = 18
    sheet["A1"].font.b = False
    sheet["A1"].fill.start_color.rgb = "FF102030"
    sheet["A1"].border.left.border_style = "double"
    sheet["A1"].border.left.color.tint = 0.5
    sheet["A1"].alignment.wrapText = True
    sheet["A1"].protection.locked = False
    output = BytesIO()
    rpyxl_wb.save(output)
    output.seek(0)
    saved = openpyxl.load_workbook(output)
    try:
        a, b = saved.active["A1"], saved.active["B1"]
        assert (a.value, b.value) == (12.5, 37.25)
        assert a.number_format == b.number_format == '0.0000 "units"'
        assert not a.font.bold and b.font.bold
        assert (a.font.sz, b.font.sz) == (18, 11)
        assert (a.font.color.rgb, b.font.color.rgb) == ("FFABCDEF", "00123456")
        assert (a.fill.fgColor.rgb, b.fill.fgColor.rgb) == ("FF102030", "00654321")
        assert (a.border.left.style, b.border.left.style) == ("double", "thin")
        assert a.border.left.color.theme == b.border.left.color.theme == 2
        assert (a.border.left.color.tint, b.border.left.color.tint) == (0.5, 0.25)
        assert a.alignment.wrap_text and not b.alignment.wrap_text
        assert a.alignment.horizontal == b.alignment.horizontal == "center"
        assert not a.protection.locked and b.protection.locked
    finally:
        saved.close()


def test_workbook_default_style_getters_do_not_mark_cells_styled(rpyxl_wb):
    sheet = rpyxl_wb.active
    for coordinate in ("A1", "B1"):
        cell = sheet[coordinate]
        assert not cell.has_style
        for name in ("font", "fill", "border", "alignment", "protection"):
            assert getattr(cell, name) is not None
            assert type(getattr(sheet._metadata.cells[coordinate], name)) is FrozenStyle
            assert not cell.has_style
        assert cell.number_format == "General"
    for name in ("font", "fill", "border", "alignment", "protection"):
        assert getattr(sheet._metadata.cells["A1"], name) is getattr(
            sheet._metadata.cells["B1"], name
        )


def test_read_only_workbook_new_frozen_getter_rejects_mutations(rpyxl_wb):
    rpyxl_wb.active["A1"] = 1
    output = BytesIO()
    rpyxl_wb.save(output)
    output.seek(0)
    loaded = load_workbook(output, read_only=True)
    try:
        cell = loaded.active["A1"]
        with pytest.warns(RuntimeWarning, match="metadata for all sheets"):
            font = cell.font
        original = loaded.active._metadata.cells["A1"].font
        assert type(original) is FrozenStyle
        assert not cell.has_style
        with pytest.raises(TypeError, match="read-only workbook"):
            font.b = True
        with pytest.raises(TypeError, match="read-only workbook"):
            font.color.tint = 0.5
        assert loaded.active._metadata.cells["A1"].font is original
        assert not cell.font.bold and cell.font.color.tint == 0
    finally:
        loaded.close()


def test_loaded_original_xml_fallback_is_writable_and_isolated_between_cells(
    openpyxl_wb,
):
    source_sheet = openpyxl_wb.active
    source_font = openpyxl.styles.Font(b=True, color="FF123456")
    for coordinate in ("A1", "B1"):
        source_sheet[coordinate] = 1
        source_sheet[coordinate].font = source_font
    source = BytesIO()
    openpyxl_wb.save(source)
    with ZipFile(source) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    styles = ET.fromstring(parts["xl/styles.xml"])
    fonts = styles.find(_metadata.q("fonts"))
    ET.SubElement(fonts[1], _metadata.q("extension"), {"value": "source"})
    parts["xl/styles.xml"] = ET.tostring(styles)
    patched = BytesIO()
    with ZipFile(patched, "w") as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    patched.seek(0)
    loaded = load_workbook(patched)
    try:
        sheet = loaded.active
        a, b = sheet["A1"], sheet["B1"]
        assert type(a.font) is type(b.font) is Font
        for model in (a.font, a.font.color, b.font, b.font.color):
            assert not (vars(model).keys() & STYLE_INTERNAL_FIELDS)
        assert a.font._xml is not b.font._xml
        assert a.font.color is not b.font.color
        sheet["C1"].font = a.font
        assert type(sheet["C1"].font) is Font
        for model in (sheet["C1"].font, sheet["C1"].font.color):
            assert not (vars(model).keys() & STYLE_INTERNAL_FIELDS)
        assert sheet["C1"].font._xml is not a.font._xml
        a.font.color.rgb = "FFABCDEF"
        a.font.b = False
        a.font._xml.find(_metadata.q("extension")).set("value", "edited")
        sheet["D1"].font = a.font
        assert not sheet["D1"].font.bold
        assert sheet["D1"].font.color.rgb == "FFABCDEF"
        assert (
            sheet["D1"].font._xml.find(_metadata.q("extension")).get("value")
            == "edited"
        )
        for cell in (b, sheet["C1"]):
            assert cell.font.bold
            assert cell.font.color.rgb == "FF123456"
            assert (
                cell.font._xml.find(_metadata.q("extension")).get("value") == "source"
            )
        sheet["D1"].font.color.tint = 0.5
        assert a.font.color.tint == b.font.color.tint == 0
        output = BytesIO()
        loaded.save(output)
        with ZipFile(output) as archive:
            saved_styles = ET.fromstring(archive.read("xl/styles.xml"))
        markers = {
            node.get("value") for node in saved_styles.iter(_metadata.q("extension"))
        }
        assert markers == {"source", "edited"}
        output.seek(0)
        saved = openpyxl.load_workbook(output)
        try:
            assert saved.active["A1"].font.color.rgb == "FFABCDEF"
            assert saved.active["B1"].font.color.rgb == "FF123456"
            assert saved.active["C1"].font.color.rgb == "FF123456"
            assert saved.active["D1"].font.color.tint == 0.5
        finally:
            saved.close()
    finally:
        loaded.close()


def test_source_cache_does_not_retain_pool_and_other_pools_intern_by_value():
    pool = StylePool()
    source = Font(color="123456")
    frozen = pool.assign(source)
    other = StylePool()
    other_value = other.assign(Font(color="123456"))
    assert other.assign(source) is other_value
    assert other_value == frozen
    reference = weakref.ref(pool)
    del pool
    gc.collect()
    assert reference() is None
