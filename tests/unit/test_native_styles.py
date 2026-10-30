"""Style preparation tests plus unskipped native/save-path integration contracts."""

from copy import deepcopy
from io import BytesIO
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import openpyxl
import pytest
from rpyxl import NativeWorkbook, Workbook, _metadata, load_workbook
from rpyxl._models import (
    CellMetadata,
    Comment,
    DefinedName,
    MetadataBook,
    MetadataSheet,
)
from rpyxl._native_styles import compile_style_records, requires_metadata_overlay
from rpyxl.chart import BarChart, Reference
from rpyxl.styles import Alignment, Border, Color, Font, PatternFill, Protection, Side


@pytest.fixture
def metadata_book():
    return MetadataBook(sheets={"Sheet": MetadataSheet()})


def test_compile_deduplicates_components_and_records(metadata_book, monkeypatch):
    cell = CellMetadata(
        font=Font(bold=True, color=Color(theme=2, tint=0.25)),
        fill=PatternFill("solid", fgColor="FF123456"),
        border=Border(left=Side("thin", "FF654321")),
        alignment=Alignment(wrap_text=True),
        protection=Protection(locked=False),
    )
    sheet = metadata_book.sheets["Sheet"]
    sheet.cells = {"A1": cell, "B2": deepcopy(cell), "C3": deepcopy(cell)}
    sheet.cells["C3"].alignment.horizontal = "center"
    metadata_book.sheets["Other"] = MetadataSheet(
        title="Other", cells={"AA7": deepcopy(cell), "B1": CellMetadata()}
    )
    counts = dict.fromkeys(("font", "fill", "border", "alignment", "protection"), 0)
    original = ET.tostring
    calls = []

    def serialize(node, *args, **kwargs):
        calls.append(node)
        counts[_metadata.local(node)] += 1
        return original(node, *args, **kwargs)

    monkeypatch.setattr(ET, "tostring", serialize)
    records, assignments = compile_style_records(metadata_book)
    assert len(records) == 2
    assert all(len(record) == 5 for record in records)
    assert all(isinstance(fragment, str) for fragment in records[0])
    assert assignments == [
        ("Sheet", 1, 1, 0),
        ("Sheet", 2, 2, 0),
        ("Sheet", 3, 3, 1),
        ("Other", 7, 27, 0),
    ]
    assert counts == {
        "font": 1,
        "fill": 1,
        "border": 1,
        "alignment": 2,
        "protection": 1,
    }
    assert len(calls) == 6  # No per-cell serialization, including fingerprints.
    assert ET.fromstring(records[0][3]).get("wrapText") == "1"
    assert "number_format" not in "".join(records[0])


def test_compile_ignores_annotation_only_and_empty_cells(metadata_book):
    metadata_book.sheets["Sheet"].cells = {
        "A1": CellMetadata(),
        "B2": CellMetadata(comment=Comment("note", "author")),
        "C3": CellMetadata(hyperlink="https://example.com"),
        "D4": CellMetadata(protection=Protection()),
    }
    records, assignments = compile_style_records(metadata_book)
    assert assignments == [("Sheet", 4, 4, 0)]
    assert records[0][:4] == (None,) * 4
    assert ET.fromstring(records[0][4]).get("locked") == "1"
    assert compile_style_records(None) == ([], [])
    metadata_book.sheets["Sheet"].cells.clear()
    assert compile_style_records(metadata_book) == ([], [])


def test_compile_observes_nested_mutation_and_never_reuses_stale_cache(metadata_book):
    cell = CellMetadata(
        font=Font(color=Color(theme=1)),
        border=Border(left=Side("thin", Color(indexed=3))),
        fill=PatternFill("solid", fgColor=Color(rgb="123456")),
    )
    metadata_book.sheets["Sheet"].cells = {"A1": cell, "B1": deepcopy(cell)}
    first, _ = compile_style_records(metadata_book)
    cell.font.color.tint = 0.5
    cell.border.left.color.indexed = 5
    cell.fill.fgColor.rgb = "FFABCDEF"
    second, assignments = compile_style_records(metadata_book)
    assert len(first) == 1
    assert len(second) == 2
    assert assignments[-1][-1] == 1
    for index in (0, 1, 2):
        assert second[0][index] != first[0][index]
        assert second[1][index] == first[0][index]
    metadata_book.sheets["Sheet"].cells["A1"] = CellMetadata()
    third, assignments = compile_style_records(metadata_book)
    assert third == first
    assert assignments == [("Sheet", 1, 2, 0)]


def test_compile_original_xml_is_structural_and_preserves_unknown_fields(
    metadata_book, monkeypatch
):
    font = _metadata._read_font(
        ET.fromstring(
            '<font xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
            'custom="original"><name val="Calibri"/><extension value="one"/></font>'
        )
    )
    alignment = Alignment(horizontal="left")
    alignment._xml = ET.fromstring('<alignment custom="yes"><unknown/></alignment>')
    sheet = metadata_book.sheets["Sheet"]
    sheet.cells = {"A1": CellMetadata(font=font, alignment=alignment)}
    sheet.cells["B1"] = deepcopy(sheet.cells["A1"])
    # Attribute insertion order is not a semantic difference.
    font._xml.set("second", "2")
    copied = sheet.cells["B1"].font._xml
    copied.attrib.clear()
    copied.attrib.update(second="2", custom="original")
    original = ET.tostring
    serialized = []

    def serialize(node, *args, **kwargs):
        serialized.append(node)
        return original(node, *args, **kwargs)

    monkeypatch.setattr(ET, "tostring", serialize)
    first, _ = compile_style_records(metadata_book)
    assert len(first) == 1
    assert len(serialized) == 2
    assert (
        ET.fromstring(first[0][0]).find(_metadata.q("extension")).get("value") == "one"
    )
    assert ET.fromstring(first[0][3]).get("custom") == "yes"
    font._xml.find(_metadata.q("extension")).set("value", "two")
    second, _ = compile_style_records(metadata_book)
    assert len(second) == 2
    assert second[0][0] != second[1][0]
    assert ET.canonicalize(first[0][0]) == ET.canonicalize(second[1][0])


def test_compile_nested_original_xml_participates_in_fingerprint(metadata_book):
    cell = CellMetadata(border=Border(left=Side("thin", Color(theme=1))))
    cell.border.left._xml = ET.fromstring('<left custom="one"/>')
    cell.border.left.color._xml = ET.fromstring('<color theme="1" custom="one"/>')
    metadata_book.sheets["Sheet"].cells = {"A1": cell, "B1": deepcopy(cell)}
    assert len(compile_style_records(metadata_book)[0]) == 1
    cell.border.left.color._xml.set("custom", "two")
    # Existing serializers decide what XML is emitted; fingerprints must not
    # silently equate original snapshots just because dataclass equality does.
    assert len(compile_style_records(metadata_book)[0]) == 2


def test_overlay_style_only_and_default_getters(metadata_book):
    assert not requires_metadata_overlay(None)
    assert not requires_metadata_overlay(metadata_book)
    sheet = metadata_book.sheets["Sheet"]
    sheet.cells["A1"] = CellMetadata(
        font=Font(),
        fill=PatternFill(),
        border=Border(),
        alignment=Alignment(),
        protection=Protection(),
    )
    _ = sheet.column_dimensions["B"], sheet.row_dimensions[2]
    assert not requires_metadata_overlay(metadata_book)
    book = Workbook()
    book.active["B2"].number_format = '0.0000 "units"'
    for name in ("font", "fill", "border", "alignment", "protection"):
        getattr(book.active["A1"], name)
    _ = book.properties, book.defined_names
    assert not requires_metadata_overlay(book._book)


@pytest.mark.parametrize(
    "feature",
    [
        "merge",
        "properties",
        "defined_names",
        "comment",
        "hyperlink",
        "width",
        "height",
        "freeze",
        "filter",
        "tab_color",
        "hidden",
        "protection",
        "page_setup",
        "validation",
        "conditional",
        "table",
        "chart",
    ],
)
def test_overlay_non_style_features(metadata_book, feature):
    sheet = metadata_book.sheets["Sheet"]
    if feature == "merge":
        sheet.merged_cells.add("A1:B1")
    elif feature == "properties":
        metadata_book.properties.title = "Report"
    elif feature == "defined_names":
        metadata_book.defined_names.add(DefinedName("Data", "Sheet!$A$1"))
    elif feature == "comment":
        sheet.cells["A1"] = CellMetadata(comment=Comment("note", "author"))
    elif feature == "hyperlink":
        sheet.cells["A1"] = CellMetadata(hyperlink="https://example.com")
    elif feature == "width":
        sheet.column_dimensions["B"].width = 20
    elif feature == "height":
        sheet.row_dimensions[2].height = 30
    elif feature == "freeze":
        sheet.freeze_panes = "B2"
    elif feature == "filter":
        sheet.auto_filter.ref = "A1:B2"
    elif feature == "tab_color":
        sheet.sheet_properties.tabColor = Color(theme=1)
    elif feature == "hidden":
        sheet.sheet_state = "hidden"
    elif feature == "protection":
        sheet.protection.sheet = False
    elif feature == "page_setup":
        sheet.page_setup.orientation = "landscape"
    elif feature == "validation":
        sheet.data_validations.append(object())
    elif feature == "conditional":
        sheet.conditional_formatting.add("A1", object())
    elif feature == "table":
        sheet.tables["Data"] = object()
    elif feature == "chart":
        sheet.charts.append(object())
    assert requires_metadata_overlay(metadata_book)


@pytest.mark.parametrize(
    "target,attribute,value",
    [
        ("book", "_parts", {}),
        ("book", "_parts", {"xl/drawings/drawing1.xml": b"xml"}),
        ("book", "_content_types", ET.Element("Types")),
        ("sheet", "_xml", ET.Element("worksheet")),
        ("sheet", "_part", "xl/worksheets/sheet1.xml"),
        ("sheet", "_rels", {}),
        ("properties", "_xml", ET.Element("properties")),
        ("auto_filter", "_xml", ET.Element("autoFilter")),
        ("column", "_xml", ET.Element("col")),
        ("row", "_xml", ET.Element("row")),
        ("validations", "_xml", ET.Element("dataValidations")),
        ("sheet", "future_feature", {}),
    ],
)
def test_overlay_retains_snapshots_and_unknown_metadata(
    metadata_book, target, attribute, value
):
    sheet = metadata_book.sheets["Sheet"]
    targets = {
        "book": metadata_book,
        "sheet": sheet,
        "properties": metadata_book.properties,
        "auto_filter": sheet.auto_filter,
        "column": sheet.column_dimensions["A"],
        "row": sheet.row_dimensions[1],
        "validations": sheet.data_validations,
    }
    setattr(targets[target], attribute, value)
    assert requires_metadata_overlay(metadata_book)


def _forbid_overlay(monkeypatch):
    def unexpected(*args, **kwargs):
        raise AssertionError("Style-only saves must bypass apply_metadata")

    monkeypatch.setattr(_metadata, "apply_metadata", unexpected)


def _read_saved(book):
    output = BytesIO()
    book.save(output)
    output.seek(0)
    return openpyxl.load_workbook(output)


def test_style_only_save_native_integration(monkeypatch):
    _forbid_overlay(monkeypatch)
    book = Workbook()
    sheet = book.active
    sheet["A1"] = 12.5
    sheet["A1"].number_format = '0.0000 "units"'
    sheet["A1"].font = Font(bold=True, color="123456")
    sheet["A1"].fill = PatternFill(
        "solid", fgColor="80112233", bgColor=Color(indexed=5)
    )
    sheet["A1"].border = Border(left=Side("thin", Color(theme=2, tint=0.25)))
    sheet["A1"].alignment = Alignment(wrap_text=True, horizontal="center")
    sheet["A1"].protection = Protection(locked=False, hidden=True)
    sheet["Z20"].font = Font(italic=True, color=Color(auto=True))
    sheet["Z20"].number_format = "0.00000"
    _ = sheet["B2"].font, sheet.column_dimensions["C"], sheet.row_dimensions[3]
    result = _read_saved(book).active
    cell = result["A1"]
    assert cell.value == 12.5
    assert cell.number_format == '0.0000 "units"'
    assert cell.font.bold and cell.font.color.rgb == "00123456"
    assert cell.fill.fgColor.rgb == "80112233"
    assert cell.fill.bgColor.indexed == 5
    assert cell.border.left.style == "thin"
    assert cell.border.left.color.theme == 2
    assert cell.border.left.color.tint == 0.25
    assert cell.alignment.wrap_text and cell.alignment.horizontal == "center"
    assert not cell.protection.locked and cell.protection.hidden
    assert result["Z20"].value is None
    assert result["Z20"].font.italic and result["Z20"].font.color.auto
    assert result["Z20"].number_format == "0.00000"


def test_style_color_assignment_copy_isolation(monkeypatch):
    _forbid_overlay(monkeypatch)
    book = Workbook()
    sheet = book.active
    font = Font(color=Color(rgb="123456"))
    border = Border(left=Side("thin", Color(theme=1)))
    for coordinate in ("A1", "B1"):
        sheet[coordinate].font = font
        sheet[coordinate].border = border
    font.color.rgb = "FF000000"
    border.left.color.theme = 5
    sheet["A1"].font.color.rgb = "FFABCDEF"
    sheet["A1"].border.left.color.theme = 3
    assert sheet["B1"].font.color.rgb == "00123456"
    assert sheet["B1"].border.left.color.theme == 1
    result = _read_saved(book).active
    assert result["A1"].font.color.rgb == "FFABCDEF"
    assert result["B1"].font.color.rgb == "00123456"
    assert result["A1"].border.left.color.theme == 3
    assert result["B1"].border.left.color.theme == 1


def test_repeated_style_saves_clear_and_reassign_native_integration(monkeypatch):
    _forbid_overlay(monkeypatch)
    book = Workbook()
    sheet = book.active
    sheet["A1"] = 3.5
    sheet["A1"].number_format = "0.0000"
    sheet["A1"].font = Font(bold=True, color="123456")
    assert _read_saved(book).active["A1"].font.bold
    sheet["A1"].font.color.rgb = "FFABCDEF"
    assert _read_saved(book).active["A1"].font.color.rgb == "FFABCDEF"
    sheet["A1"].font = None
    cleared = _read_saved(book).active["A1"]
    assert not cleared.font.bold
    assert cleared.number_format == "0.0000"
    sheet["B2"].fill = PatternFill("solid", fgColor="FF102030")
    reassigned = _read_saved(book).active
    assert not reassigned["A1"].font.bold
    assert reassigned["A1"].number_format == "0.0000"
    assert reassigned["B2"].fill.fgColor.rgb == "FF102030"
    sheet["B2"].fill = None
    assert _read_saved(book).active["B2"].fill.patternType is None


def test_metadata_fallback_preserves_charts_comments_and_loaded_content(monkeypatch):
    calls = []
    original = _metadata.apply_metadata

    def overlay(*args, **kwargs):
        calls.append(True)
        return original(*args, **kwargs)

    monkeypatch.setattr(_metadata, "apply_metadata", overlay)
    book = Workbook()
    sheet = book.active
    sheet.append(["Amount", 1])
    sheet.append(["Value", 2])
    sheet["B2"].font = Font(bold=True)
    sheet["B2"].comment = Comment("Keep this note", "Reviewer")
    sheet.merge_cells("D1:E1")
    book.properties.title = "Keep this title"
    chart = BarChart()
    chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=2))
    sheet.add_chart(chart, "G1")
    source = BytesIO()
    book.save(source)
    source.seek(0)
    loaded = load_workbook(source)
    assert requires_metadata_overlay(loaded._book)
    result = _read_saved(loaded)
    assert len(calls) == 2
    assert result.properties.title == "Keep this title"
    assert result.active["B2"].comment.text == "Keep this note"
    assert result.active["B2"].font.bold
    assert len(result.active._charts) == 1
    assert "D1:E1" in result.active.merged_cells


def test_overlay_retains_native_styles_without_rebuilding_them(monkeypatch):
    book = NativeWorkbook()
    book.active["A1"] = 12.5
    book.active["A1"].number_format = "0.0000"
    generated = BytesIO()
    book.save(generated)
    metadata = MetadataBook(sheets={"Sheet": MetadataSheet()})
    metadata.sheets["Sheet"].cells["A1"] = CellMetadata(font=Font(bold=True))
    styled_source = BytesIO()
    _metadata.apply_metadata(generated, styled_source, metadata)

    # Simulate a native-generated styled archive; other features must leave its
    # style IDs and number formats intact instead of overlaying cell styles again.
    metadata.sheets["Sheet"].cells["A1"].font = Font(italic=True)
    metadata.sheets["Sheet"].cells["A1"].comment = Comment("Keep styles", "Reviewer")

    def unexpected(*args, **kwargs):
        raise AssertionError(
            "Cell styles already serialized by Rust must not be rebuilt"
        )

    monkeypatch.setattr(_metadata._Styles, "overlay", unexpected)
    output = BytesIO()
    _metadata.apply_metadata(styled_source, output, metadata, native_styles=True)
    loaded = openpyxl.load_workbook(output)
    try:
        cell = loaded.active["A1"]
        assert cell.value == 12.5
        assert cell.font.bold and not cell.font.italic
        assert cell.number_format == "0.0000"
        assert cell.comment.text == "Keep styles"
    finally:
        loaded.close()
        book.close()


def test_native_workbook_remains_raw():
    book = NativeWorkbook()
    book.active["A1"] = "native"
    assert not hasattr(book, "_book")
    output = BytesIO()
    book.save(output)
    output.seek(0)
    with ZipFile(output) as archive:
        assert any(
            b"native" in archive.read(part)
            for part in ("xl/sharedStrings.xml", "xl/worksheets/sheet1.xml")
            if part in archive.namelist()
        )
    output.seek(0)
    assert openpyxl.load_workbook(output).active["A1"].value == "native"
