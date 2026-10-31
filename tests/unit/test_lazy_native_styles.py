"""Source decoder tests and unskipped contracts for the new native wheel.

Integration tests deliberately assert the real API: an old installed wheel is a
failure, not a skip or a mocked implementation of the performance contract.
"""

from copy import deepcopy
from io import BytesIO
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import openpyxl
import pytest
from rpyxl import NativeWorkbook, _metadata, load_workbook
from rpyxl._compat import structural_metadata
from rpyxl._models import CellMetadata, MetadataBook, MetadataSheet
from rpyxl._native_style_read import decode_record, preserve_opaque_records
from rpyxl._native_styles import compile_style_records
from rpyxl._style_pool import FrozenStyle, StylePool
from rpyxl.cell import MergedCell
from rpyxl.styles import Font

S = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


def _xml(body):
    return body.replace("<font", '<font xmlns="' + S[1:-1] + '"', 1)


@pytest.mark.parametrize(
    "body,clean",
    [
        ('<font><name val="Calibri"/><color theme="1"/></font>', True),
        ('<font custom="yes"><name val="Calibri"/></font>', False),
        ('<font><name val="Calibri" custom="yes"/></font>', False),
        ('<font><color theme="1" custom="yes"/></font>', False),
        ('<font><name val="Calibri"/><name val="Other"/></font>', False),
        ('<font><foreign xmlns="urn:extension"/></font>', False),
        ('<font><name val="Calibri">unmodeled</name></font>', False),
    ],
)
def test_decode_strips_only_strictly_modeled_xml(body, clean):
    value = decode_record((_xml(body), None, None, None, None), StylePool())[0]
    assert (type(value) is FrozenStyle) is clean
    if clean:
        assert type(value.color) is FrozenStyle
        assert "_xml" not in vars(value)
        assert "_xml" not in vars(value.color)
    else:
        assert value._xml is not None


def test_decode_gradient_remains_opaque_and_known_alignment_freezes():
    record = (
        None,
        f'<fill xmlns="{S[1:-1]}"><gradientFill degree="45"><stop position="0"><color rgb="FFFF0000"/></stop></gradientFill></fill>',
        None,
        f'<alignment xmlns="{S[1:-1]}" wrapText="1" textRotation="15"/>',
        None,
    )
    values = decode_record(record, StylePool())
    assert type(values[1]) is not FrozenStyle
    assert values[1]._xml.find(S + "gradientFill") is not None
    assert type(values[3]) is FrozenStyle
    assert values[3].wrap_text
    assert values[3].text_rotation == 15


def test_opaque_serialization_keeps_nested_extension_attributes():
    font = decode_record(
        (
            _xml(
                '<font custom="root"><name val="Calibri" producer="name"/><color theme="1" producer="color"/><ext xmlns="urn:extension"/></font>'
            ),
            None,
            None,
            None,
            None,
        ),
        StylePool(),
    )[0]
    font.bold = True
    book = MetadataBook(sheets={"Sheet": MetadataSheet(title="Sheet")})
    book.sheets["Sheet"].cells["A1"] = CellMetadata(font=font, style_dirty=True)
    records, assignments = compile_style_records(book)
    records = preserve_opaque_records(book, records, assignments)
    node = ET.fromstring(records[0][0])
    assert node.get("custom") == "root"
    assert node.find(S + "name").get("producer") == "name"
    assert node.find(S + "color").get("producer") == "color"
    assert node.find("{urn:extension}ext") is not None
    assert node.find(S + "b") is not None


def _source():
    book = openpyxl.Workbook()
    sheet = book.active
    for coordinate in ("A1", "B1", "C1"):
        sheet[coordinate] = coordinate
        sheet[coordinate].font = openpyxl.styles.Font(
            name="Source Font", bold=True, color="FF123456"
        )
        sheet[coordinate].fill = openpyxl.styles.PatternFill(
            "solid", fgColor="FFABCDEF"
        )
        sheet[coordinate].alignment = openpyxl.styles.Alignment(horizontal="center")
    sheet["B1"].comment = openpyxl.comments.Comment("Original comment", "Author")
    sheet["B1"].hyperlink = "https://example.com/original"
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def _rewrite(data, part, edit):
    output = BytesIO()
    with ZipFile(BytesIO(data)) as source, ZipFile(output, "w") as target:
        for name in source.namelist():
            payload = source.read(name)
            target.writestr(name, edit(payload) if name == part else payload)
    return output.getvalue()


def _load(data):
    for name in (
        "_has_original_package",
        "_can_preserve_package",
        "_original_package",
        "_source_style_record",
        "_update_cell_styles",
    ):
        assert hasattr(NativeWorkbook, name), (
            f"Rebuild the native wheel: missing {name}"
        )
    book = load_workbook(BytesIO(data))
    assert book._native_package_api
    assert book._native_style_reader is not None
    assert book._native._has_original_package()
    return book


def _saved(book):
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def _forbid(*args, **kwargs):
    raise AssertionError(
        "Ordinary native load/style access must not parse Python metadata"
    )


def test_load_and_style_get_never_load_metadata_or_scan_worksheets(monkeypatch):
    data = _source()
    monkeypatch.setattr(_metadata, "load_metadata", _forbid)
    monkeypatch.setattr(_metadata._Styles, "read", _forbid)
    monkeypatch.setattr(_metadata, "_worksheet_metadata_xml", _forbid)
    monkeypatch.setattr(_metadata, "apply_metadata", _forbid)
    book = _load(data)
    try:
        assert book._book is None
        assert book.active["A1"].font.name == "Source Font"
        assert book.active["B1"].font.bold
        assert book.active["A1"].has_style
        assert book._book is None
        assert len(book._native_style_reader.records) == 1
        _saved(book)
        assert book._book is None
    finally:
        book.close()


def test_sparse_edit_preserves_untouched_styles_and_nested_cow(monkeypatch):
    book = _load(_source())
    monkeypatch.setattr(_metadata, "load_metadata", _forbid)
    monkeypatch.setattr(_metadata, "apply_metadata", _forbid)
    try:
        before = deepcopy(book.active["B1"].font)
        book.active["A1"].font.color.rgb = "FF987654"
        assert book.active["B1"].font == before
        assert set(book._book.sheets["Sheet"].cells) == {"A1"}
        assert not book._full_metadata_loaded
        result = openpyxl.load_workbook(BytesIO(_saved(book)))
        assert result.active["A1"].font.color.rgb == "FF987654"
        for coordinate in ("B1", "C1"):
            assert result.active[coordinate].font.name == "Source Font"
            assert result.active[coordinate].font.bold
            assert result.active[coordinate].fill.fgColor.rgb == "FFABCDEF"
        assert result.active["A1"].alignment.horizontal == "center"
        assert result.active["B1"].comment.text == "Original comment"
        assert result.active["B1"].hyperlink.target == "https://example.com/original"
    finally:
        book.close()


def test_repeated_clear_is_sparse_and_explicit():
    book = _load(_source())
    try:
        for _ in range(2):
            for name in ("font", "fill", "border", "alignment", "protection"):
                setattr(book.active["A1"], name, None)
            result = openpyxl.load_workbook(BytesIO(_saved(book)))
            assert not result.active["A1"].font.bold
            assert result.active["A1"].fill.patternType is None
            assert result.active["B1"].font.bold
        book.active["A1"].font = Font(italic=True)
        result = openpyxl.load_workbook(BytesIO(_saved(book)))
        assert result.active["A1"].font.italic
        assert result.active["B1"].font.bold
    finally:
        book.close()


def _opaque_source():
    def edit(data):
        root = ET.fromstring(data)
        for node in root.findall(S + "fonts/" + S + "font"):
            node.set("producer", "root")
            node.find(S + "color").set("producer", "color")
        return ET.tostring(root)

    return _rewrite(_source(), "xl/styles.xml", edit)


def test_opaque_get_is_copy_on_first_cell_and_does_not_rewrite(monkeypatch):
    data = _opaque_source()
    book = _load(data)
    monkeypatch.setattr(_metadata, "load_metadata", _forbid)
    try:
        first = book.active["A1"].font
        assert first is book.active["A1"].font
        assert first is not book.active["B1"].font
        assert book._book is None
        saved = _saved(book)
        assert book._book is None
        with ZipFile(BytesIO(data)) as original, ZipFile(BytesIO(saved)) as output:
            assert original.read("xl/styles.xml") == output.read("xl/styles.xml")
        first.italic = True
        saved = _saved(book)
        assert set(book._book.sheets["Sheet"].cells) == {"A1"}
        with _load(saved) as result:
            assert result.active["A1"].font.italic
            assert not result.active["B1"].font.italic
        with ZipFile(BytesIO(saved)) as output:
            root = ET.fromstring(output.read("xl/styles.xml"))
            fonts = root.findall(S + "fonts/" + S + "font")
            edited = next(node for node in fonts if node.find(S + "i") is not None)
            assert edited.get("producer") == "root"
            assert edited.find(S + "color").get("producer") == "color"
    finally:
        book.close()


def test_full_metadata_promotion_merges_sparse_styles_and_annotations(monkeypatch):
    book = _load(_source())
    original = _metadata.load_metadata
    calls = []

    def load(source, *, cell_styles=True):
        calls.append(cell_styles)
        return original(source, cell_styles=cell_styles)

    monkeypatch.setattr(_metadata, "load_metadata", load)
    monkeypatch.setattr(_metadata._Styles, "read", _forbid)
    try:
        book.active["B1"].font.italic = True
        book.properties.title = "Rich edit"
        assert calls == [False]
        assert book.active["B1"].comment.text == "Original comment"
        assert book.active["B1"].hyperlink.target == "https://example.com/original"
        result = openpyxl.load_workbook(BytesIO(_saved(book)))
        assert result.properties.title == "Rich edit"
        assert result.active["B1"].font.italic
        assert result.active["B1"].comment.text == "Original comment"
        assert result.active["C1"].font.bold
        assert calls == [False]
    finally:
        book.close()


def test_original_xml_features_survive_style_only_native_save(monkeypatch):
    def edit(data):
        root = ET.fromstring(data)
        ET.SubElement(root, S + "extLst").append(
            ET.Element("{urn:producer}feature", {"value": "keep"})
        )
        return ET.tostring(root)

    data = _rewrite(_source(), "xl/worksheets/sheet1.xml", edit)
    book = _load(data)
    monkeypatch.setattr(_metadata, "load_metadata", _forbid)
    monkeypatch.setattr(_metadata, "apply_metadata", _forbid)
    try:
        book.active["A1"].font.italic = True
        with (
            ZipFile(BytesIO(data)) as original,
            ZipFile(BytesIO(_saved(book))) as output,
        ):
            root = ET.fromstring(output.read("xl/worksheets/sheet1.xml"))
            assert root.find(S + "extLst/{urn:producer}feature").get("value") == "keep"
            assert original.read("xl/theme/theme1.xml") == output.read(
                "xl/theme/theme1.xml"
            )
            assert original.read("xl/_rels/workbook.xml.rels") == output.read(
                "xl/_rels/workbook.xml.rels"
            )
    finally:
        book.close()


def test_unsupported_template_retains_stable_snapshot_and_falls_back():
    # Numeric escapes conservatively disqualify the native raw-copy template
    # (they can hide macro content types), but remain valid ordinary OOXML.
    data = _rewrite(
        _source(),
        "[Content_Types].xml",
        lambda value: value.replace(b'Extension="xml"', b'Extension="&#120;ml"'),
    )
    book = _load(data)
    try:
        assert not book._native._can_preserve_package()
        assert book._native._original_package() == data
        assert book._book is None
        book.active["A1"].font.italic = True
        result = openpyxl.load_workbook(BytesIO(_saved(book)))
        assert book._full_metadata_loaded
        assert result.active["A1"].font.italic
        assert result.active["C1"].font.bold
        assert result.active["B1"].comment.text == "Original comment"
    finally:
        book.close()


def test_loaded_structure_checks_original_names_before_changes():
    source = openpyxl.load_workbook(BytesIO(_source()))
    source.defined_names.add(
        openpyxl.workbook.defined_name.DefinedName("Guard", attr_text="Sheet!$A$1")
    )
    output = BytesIO()
    source.save(output)
    book = _load(output.getvalue())
    try:
        assert book._book is None
        with pytest.raises(NotImplementedError, match="defined names"):
            book.active.insert_rows(1)
        with pytest.raises(NotImplementedError, match="names"):
            book.active.title = "Renamed"
        assert book.sheetnames == ["Sheet"]
    finally:
        book.close()


def _merged_source():
    source = openpyxl.Workbook()
    sheet = source.active
    sheet["A1"] = "outside"
    sheet["B4"] = "anchor"
    sheet["E7"] = "after"
    sheet.merge_cells("B4:D6")
    for key in ("A1", "B4", "C4", "D6"):
        cell = sheet[key]
        cell.font = openpyxl.styles.Font(name="Merged Source", bold=True)
        cell.fill = openpyxl.styles.PatternFill("solid", fgColor="FFABCDEF")
        cell.border = openpyxl.styles.Border(
            bottom=openpyxl.styles.Side(style="thin", color="FF123456")
        )
        cell.alignment = openpyxl.styles.Alignment(horizontal="center")
        cell.protection = openpyxl.styles.Protection(locked=False)
        cell.number_format = "0.000"
    output = BytesIO()
    source.save(output)
    return output.getvalue()


def _merged_metadata_state(book, state):
    assert book._book is None
    if state == "sparse":
        book.active["A1"].font.italic = True
        assert set(book._book.sheets["Sheet"].cells) == {"A1"}
    elif state == "full":
        assert {str(value) for value in book.active.merged_cells} == {"B4:D6"}
        assert not book._book.sheets["Sheet"].cells
    assert book._full_metadata_loaded is (state == "full")


@pytest.mark.parametrize("state", ["lazy", "sparse", "full"])
@pytest.mark.parametrize("keyword_bounds", [False, True], ids=["range", "bounds"])
def test_loaded_unmerge_clears_unaccessed_native_styles(
    monkeypatch, state, keyword_bounds
):
    book = _load(_merged_source())
    monkeypatch.setattr(_metadata._Styles, "read", _forbid)
    try:
        _merged_metadata_state(book, state)
        sheet = book.active
        if keyword_bounds:
            sheet.unmerge_cells(start_row=4, start_column=2, end_row=6, end_column=4)
        else:
            sheet.unmerge_cells("B4:D6")
        assert not sheet.merged_cells.ranges
        cells = book._book.sheets["Sheet"].cells
        # Only source-styled descendants need overrides, not the merge rectangle
        # or the other styled cells whose native records must remain untouched.
        assert set(cells) == ({"A1"} if state == "sparse" else set()) | {"C4", "D6"}
        for key in ("C4", "D6"):
            assert cells[key] == CellMetadata(style_dirty=True)
        records, assignments = compile_style_records(book._book)
        for row, col in ((4, 3), (6, 4)):
            index = next(
                index
                for title, r, c, index in assignments
                if (title, r, c) == ("Sheet", row, col)
            )
            assert records[index] == (None,) * 5
        descendants = ("C4", "D4", "B5", "C5", "D5", "B6", "C6", "D6")
        for key in descendants:
            assert not isinstance(sheet[key], MergedCell)
            assert sheet[key].value is None
            assert not sheet[key].has_style
            assert sheet[key].number_format == "General"
        assert sheet["B4"].value == "anchor"
        assert sheet["B4"].font.bold
        assert sheet["B4"].fill.fgColor.rgb == "FFABCDEF"
        assert sheet["B4"].border.bottom.style == "thin"
        assert sheet["B4"].alignment.horizontal == "center"
        assert not sheet["B4"].protection.locked
        assert sheet["B4"].number_format == "0.000"
        assert sheet["A1"].font.bold
        sheet["C5"] = "now writable"
        for _ in range(2):
            result = openpyxl.load_workbook(BytesIO(_saved(book)))
            try:
                assert not result.active.merged_cells.ranges
                assert result.active["C5"].value == "now writable"
                for key in descendants:
                    assert not result.active[key].has_style
                assert result.active["B4"].value == "anchor"
                assert result.active["B4"].font.bold
                assert result.active["B4"].fill.fgColor.rgb == "FFABCDEF"
                assert result.active["B4"].border.bottom.style == "thin"
                assert result.active["B4"].alignment.horizontal == "center"
                assert not result.active["B4"].protection.locked
                assert result.active["B4"].number_format == "0.000"
                assert result.active["A1"].font.bold
                assert result.active["A1"].font.italic is (state == "sparse")
                assert result.active["E7"].value == "after"
            finally:
                result.close()
    finally:
        book.close()


@pytest.mark.parametrize("state", ["lazy", "sparse", "full"])
@pytest.mark.parametrize(
    "access",
    ["setitem", "index", "range", "cell", "cell_value", "rows", "cols", "iter"],
)
def test_loaded_merge_rejects_descendant_writes_without_metadata_load(
    monkeypatch, state, access
):
    book = _load(_merged_source())
    try:
        _merged_metadata_state(book, state)
        monkeypatch.setattr(_metadata, "load_metadata", _forbid)
        sheet = book.active
        before = book._book
        before_cells = deepcopy(before.sheets["Sheet"].cells) if before else None
        before_dimensions = sheet.dimensions
        with pytest.raises(AttributeError, match="MergedCell.*read-only"):
            if access == "setitem":
                sheet["C5"] = "rejected"
            elif access == "cell_value":
                sheet.cell(row=5, column=3, value="rejected")
            else:
                if access == "index":
                    cell = sheet["C5"]
                elif access == "range":
                    cell = sheet["C5:D6"][0][0]
                elif access == "cell":
                    cell = sheet.cell(row=5, column=3)
                elif access == "rows":
                    cell = next(
                        sheet.iter_rows(min_row=5, max_row=5, min_col=3, max_col=3)
                    )[0]
                elif access == "cols":
                    cell = next(
                        sheet.iter_cols(min_row=5, max_row=5, min_col=3, max_col=3)
                    )[0]
                else:
                    cell = list(sheet)[4][2]
                assert isinstance(cell, MergedCell)
                cell.value = "rejected"
        assert sheet._native["C5"].value is None
        assert sheet["B4"].value == "anchor"
        assert sheet.dimensions == before_dimensions
        assert book._book is before
        if before is not None:
            assert before.sheets["Sheet"].cells == before_cells
        assert book._full_metadata_loaded is (state == "full")
        # Anchor and nearby coordinates outside the off-diagonal range stay
        # writable; a swapped native row/column test must not over-guard them.
        for key in ("B4", "C3", "E3", "A5", "E5", "C7"):
            sheet[key] = "accepted"
            assert sheet[key].value == "accepted"
    finally:
        book.close()


@pytest.mark.parametrize("state", ["lazy", "sparse", "full"])
@pytest.mark.parametrize(
    "operation,args,kwargs",
    [
        ("insert_rows", (4,), {}),
        ("delete_rows", (4,), {}),
        ("insert_cols", (2,), {}),
        ("delete_cols", (2,), {}),
        ("move_range", ("B4:D6",), {"rows": 2, "cols": 1}),
    ],
)
def test_loaded_merge_guards_structure_without_changing_sheet(
    state, operation, args, kwargs
):
    book = _load(_merged_source())
    try:
        _merged_metadata_state(book, state)
        sheet = book.active
        before_values = tuple(sheet.values)
        before_dimensions = sheet.dimensions
        before_style = deepcopy(sheet["B4"].font)
        before_changed = book._structure_changed
        with pytest.raises(NotImplementedError, match="merges"):
            getattr(sheet, operation)(*args, **kwargs)
        assert book.sheetnames == ["Sheet"]
        assert tuple(sheet.values) == before_values
        assert sheet.dimensions == before_dimensions
        assert sheet["B4"].font == before_style
        assert sheet["B4"].number_format == "0.000"
        assert book._structure_changed is before_changed
        assert {str(value) for value in sheet.merged_cells} == {"B4:D6"}
        assert set(book._book.sheets["Sheet"].cells) == (
            {"A1"} if state == "sparse" else set()
        )
        assert isinstance(sheet["C5"], MergedCell)
        result = openpyxl.load_workbook(BytesIO(_saved(book)))
        try:
            assert {str(value) for value in result.active.merged_cells} == {"B4:D6"}
            assert result.active["B4"].value == "anchor"
            assert result.active["A1"].value == "outside"
            assert result.active["E7"].value == "after"
            assert result.active["B4"].font.bold
            assert result.active["B4"].number_format == "0.000"
            assert result.active["A1"].font.italic is (state == "sparse")
        finally:
            result.close()
    finally:
        book.close()


def test_native_copy_keeps_source_styles_with_only_sparse_metadata():
    book = _load(_source())
    try:
        book.active["A1"].font.italic = True
        copied = book.copy_worksheet(book.active)
        # Full feature metadata includes annotations, not an original style grid.
        assert set(book._book.sheets[copied.title].cells) == {"A1", "B1"}
        assert book._book.sheets[copied.title].cells["B1"].font is None
        assert copied["C1"].font.bold
        copied["A1"].font.bold = False
        assert book.active["A1"].font.bold
        result = openpyxl.load_workbook(BytesIO(_saved(book)))
        assert result[copied.title]["A1"].font.italic
        assert not result[copied.title]["A1"].font.bold
        assert result[copied.title]["C1"].font.bold
        assert result[copied.title]["B1"].comment.text == "Original comment"
    finally:
        book.close()


def test_sheet_insertion_updates_unloaded_local_name_scopes():
    source = openpyxl.load_workbook(BytesIO(_source()))
    source.active.defined_names.add(
        openpyxl.workbook.defined_name.DefinedName(
            "Local", attr_text="Sheet!$A$1", localSheetId=0
        )
    )
    output = BytesIO()
    source.save(output)
    book = _load(output.getvalue())
    try:
        assert book._book is None
        book.create_sheet("First", index=0)
        assert book._full_metadata_loaded
        assert book.defined_names["Local"].localSheetId == 1
        result = openpyxl.load_workbook(BytesIO(_saved(book)))
        assert "Local" in result["Sheet"].defined_names
        assert "Local" not in result["First"].defined_names
    finally:
        book.close()


def _dimension_source():
    book = openpyxl.Workbook()
    sheet = book.active
    sheet["A1"] = "Dimension defaults"
    row = sheet.row_dimensions[2]
    row.font = openpyxl.styles.Font(name="Row Default", bold=True)
    row.number_format = "0.000"
    column = sheet.column_dimensions["D"]
    column.font = openpyxl.styles.Font(name="Column Default", italic=True)
    column.number_format = "0.0000"
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def _without_generated_dimensions(data):
    def edit(payload):
        root = ET.fromstring(payload)
        for row in root.findall(S + "sheetData/" + S + "row"):
            index = row.get("r")
            row.attrib.clear()
            row.set("r", index)
        for cols in root.findall(S + "cols"):
            root.remove(cols)
        return ET.tostring(root)

    return _rewrite(data, "xl/worksheets/sheet1.xml", edit)


def _assert_dimension_styles(original, saved, *, mutated=False):
    source = openpyxl.load_workbook(BytesIO(original))
    result = openpyxl.load_workbook(BytesIO(saved))
    try:
        row = result.active.row_dimensions[2]
        column = result.active.column_dimensions["D"]
        assert result.properties.title == "Dimension overlay"
        assert row.style_id == source.active.row_dimensions[2].style_id
        assert column.style_id == source.active.column_dimensions["D"].style_id
        assert row.font.name == "Row Default"
        assert row.font.bold
        assert row.number_format == "0.000"
        assert column.font.name == "Column Default"
        assert column.font.italic
        assert column.number_format == "0.0000"
        if mutated:
            assert row.height == 24
            assert row.hidden
            assert column.width == 19
            assert column.hidden
        # openpyxl can re-intern custom number formats while reading, so compare
        # raw XF references against the source as well as the decoded defaults.
        with ZipFile(BytesIO(original)) as before, ZipFile(BytesIO(saved)) as after:
            original_root = ET.fromstring(before.read("xl/worksheets/sheet1.xml"))
            root = ET.fromstring(after.read("xl/worksheets/sheet1.xml"))
            path = S + "sheetData/" + S + 'row[@r="2"]'
            row_node = root.find(path)
            assert row_node.get("s") == original_root.find(path).get("s")
            assert row_node.get("customFormat") == "1"
            path = S + "cols/" + S + 'col[@min="4"]'
            assert root.find(path).get("style") == original_root.find(path).get("style")
    finally:
        source.close()
        result.close()


@pytest.mark.parametrize("mutated", [False, True])
def test_generated_overlay_preserves_source_dimension_styles(mutated):
    data = _dimension_source()
    metadata = _metadata.load_metadata(BytesIO(data), cell_styles=False)
    metadata.properties.title = "Dimension overlay"
    if mutated:
        row = metadata.sheets["Sheet"].row_dimensions[2]
        row.height = 24
        row.hidden = True
        column = metadata.sheets["Sheet"].column_dimensions["D"]
        column.width = 19
        column.hidden = True
    output = BytesIO()
    _metadata.apply_metadata(
        BytesIO(_without_generated_dimensions(data)),
        output,
        metadata,
        native_styles=True,
    )
    _assert_dimension_styles(data, output.getvalue(), mutated=mutated)


def test_generated_overlay_retains_unmodeled_row_attrs_at_shifted_index():
    data = _dimension_source()

    def edit(payload):
        root = ET.fromstring(payload)
        row = root.find(S + "sheetData/" + S + 'row[@r="2"]')
        row.set("thickTop", "1")
        return ET.tostring(root)

    data = _rewrite(data, "xl/worksheets/sheet1.xml", edit)
    metadata = _metadata.load_metadata(BytesIO(data), cell_styles=False)
    sheet = metadata.sheets["Sheet"]
    source_row = sheet.row_dimensions[2]
    assert source_row._xml.get("s") is not None
    metadata.sheets["Sheet"] = structural_metadata(sheet, "row", 1, 1)
    output = BytesIO()
    _metadata.apply_metadata(
        BytesIO(_without_generated_dimensions(data)),
        output,
        metadata,
        native_styles=True,
    )
    with ZipFile(output) as archive:
        root = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
        row = root.find(S + "sheetData/" + S + 'row[@r="3"]')
        assert row is not None
        assert row.get("s") == source_row._xml.get("s")
        assert row.get("customFormat") == "1"
        assert row.get("thickTop") == "1"
        old_row = root.find(S + "sheetData/" + S + 'row[@r="2"]')
        assert old_row is None or old_row.get("s") is None


@pytest.mark.parametrize("mutated", [False, True])
def test_native_rich_edit_preserves_source_dimension_styles(mutated):
    data = _dimension_source()
    with _load(data) as book:
        assert book._book is None
        book.properties.title = "Dimension overlay"
        if mutated:
            book.active.row_dimensions[2].height = 24
            book.active.row_dimensions[2].hidden = True
            book.active.column_dimensions["D"].width = 19
            book.active.column_dimensions["D"].hidden = True
        _assert_dimension_styles(data, _saved(book), mutated=mutated)


@pytest.mark.parametrize("operation", ["create", "rename", "create_then_rename"])
def test_held_opaque_font_survives_sheet_structure_changes(operation):
    with _load(_opaque_source()) as book:
        sheet = book.active
        held = sheet["A1"].font
        assert book._book is None
        if operation in ("create", "create_then_rename"):
            book.create_sheet("Other")
            assert sheet["A1"].font is held
        if operation in ("rename", "create_then_rename"):
            sheet.title = "Renamed"
            assert sheet["A1"].font is held
        metadata = book._book.sheets[sheet.title].cells["A1"]
        assert metadata.font is held
        held.italic = True
        assert sheet["A1"].font.italic
        with _load(_saved(book)) as result:
            assert result[sheet.title]["A1"].font.italic
            assert not result[sheet.title]["B1"].font.italic
            assert result[sheet.title]["A1"].font._xml.get("producer") == "root"
