import os
import openpyxl
import rpyxl
import pytest

FIXTURES_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "fixtures")

@pytest.mark.parametrize("filename", ["sample_types.xlsx", "multilingual.xlsx", "multi_sheet.xlsx", "sparse.xlsx"])
def test_compat_sheetnames_and_active(filename):
    path = os.path.join(FIXTURES_DIR, filename)

    op_wb = openpyxl.load_workbook(path)
    px_wb = rpyxl.load_workbook(path)

    assert px_wb.sheetnames == op_wb.sheetnames
    assert px_wb.active.title == op_wb.active.title
    assert len(px_wb.sheetnames) == len(op_wb.sheetnames)


    op_wb.close()
    px_wb.close()

def test_compat_sample_types_cell_values():
    path = os.path.join(FIXTURES_DIR, "sample_types.xlsx")

    op_wb = openpyxl.load_workbook(path)
    px_wb = rpyxl.load_workbook(path)

    op_ws = op_wb.active
    px_ws = px_wb.active

    for r in range(1, 6):
        for c in range(1, 6):
            op_val = op_ws.cell(r, c).value
            px_val = px_ws.cell(r, c).value
            assert px_val == op_val, f"Mismatch at ({r}, {c}): rpyxl={px_val!r} vs openpyxl={op_val!r}"

    op_wb.close()
    px_wb.close()

def test_compat_multilingual_values():
    path = os.path.join(FIXTURES_DIR, "multilingual.xlsx")

    op_wb = openpyxl.load_workbook(path)
    px_wb = rpyxl.load_workbook(path)

    op_rows = list(op_wb.active.iter_rows(values_only=True))
    px_rows = list(px_wb.active.iter_rows(values_only=True))

    assert len(px_rows) == len(op_rows)
    for i, (op_row, px_row) in enumerate(zip(op_rows, px_rows)):
        assert px_row == op_row, f"Row {i} mismatch: rpyxl={px_row} vs openpyxl={op_row}"

    op_wb.close()
    px_wb.close()

def test_compat_read_only_streaming():
    path = os.path.join(FIXTURES_DIR, "sample_types.xlsx")

    op_wb = openpyxl.load_workbook(path, read_only=True)
    px_wb = rpyxl.load_workbook(path, read_only=True)

    op_ws = op_wb.active
    px_ws = px_wb.active

    assert px_ws.title == op_ws.title
    assert px_ws.min_row == op_ws.min_row
    assert px_ws.max_row == op_ws.max_row
    assert px_ws.min_column == op_ws.min_column
    assert px_ws.max_column == op_ws.max_column

    op_rows = list(op_ws.iter_rows(values_only=True))
    px_rows = list(px_ws.iter_rows(values_only=True))

    assert len(px_rows) == len(op_rows)
    for i, (op_r, px_r) in enumerate(zip(op_rows, px_rows)):
        assert px_r == op_r, f"Row {i} mismatch in read_only mode"

    op_wb.close()
    px_wb.close()

def test_compat_sparse_sheet_iteration():
    path = os.path.join(FIXTURES_DIR, "sparse.xlsx")

    op_wb = openpyxl.load_workbook(path, read_only=True)
    px_wb = rpyxl.load_workbook(path, read_only=True)

    op_rows = list(op_wb.active.iter_rows(values_only=True))
    px_rows = list(px_wb.active.iter_rows(values_only=True))

    assert len(px_rows) == len(op_rows)
    for i in range(len(op_rows)):
        assert px_rows[i] == op_rows[i], f"Sparse row {i} mismatch: {px_rows[i]} vs {op_rows[i]}"

    op_wb.close()
    px_wb.close()
