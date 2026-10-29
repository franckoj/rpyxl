"""Compatibility test suite comparing openpyxl and RPyXL behavior directly."""

import pytest
import rpyxl
import openpyxl


def test_compat_workbook_defaults():
    rpyxl_wb = rpyxl.Workbook()
    openpyxl_wb = openpyxl.Workbook()

    assert rpyxl_wb.sheetnames == openpyxl_wb.sheetnames
    assert rpyxl_wb.active.title == openpyxl_wb.active.title
    assert len(rpyxl_wb.worksheets) == len(openpyxl_wb.worksheets)
    assert len(rpyxl_wb.sheetnames) == len(openpyxl_wb.sheetnames)

    rpyxl_wb.close()
    openpyxl_wb.close()


def test_compat_sheet_creation_and_removal():
    p_wb = rpyxl.Workbook()
    o_wb = openpyxl.Workbook()

    p_ws1 = p_wb.create_sheet(title="Custom1")
    o_ws1 = o_wb.create_sheet(title="Custom1")

    assert p_wb.sheetnames == o_wb.sheetnames
    assert p_ws1.title == o_ws1.title

    p_ws0 = p_wb.create_sheet(title="First", index=0)
    o_ws0 = o_wb.create_sheet(title="First", index=0)

    assert p_wb.sheetnames == o_wb.sheetnames
    assert p_ws0.title == o_ws0.title

    p_wb.remove(p_ws1)
    o_wb.remove(o_ws1)

    assert p_wb.sheetnames == o_wb.sheetnames

    p_wb.close()
    o_wb.close()


def test_compat_cell_types_and_values():
    p_wb = rpyxl.Workbook()
    o_wb = openpyxl.Workbook()

    p_ws = p_wb.active
    o_ws = o_wb.active

    test_data = [
        ("A1", "Text String"),
        ("A2", 12345),
        ("A3", 99.99),
        ("A4", True),
        ("A5", False),
        ("A6", None),
    ]

    for coord, val in test_data:
        p_ws[coord] = val
        o_ws[coord] = val

    for coord, _ in test_data:
        assert p_ws[coord].value == o_ws[coord].value
        assert type(p_ws[coord].value) == type(o_ws[coord].value)
        assert p_ws[coord].coordinate == o_ws[coord].coordinate
        assert p_ws[coord].row == o_ws[coord].row
        assert p_ws[coord].column == o_ws[coord].column
        assert p_ws[coord].column_letter == o_ws[coord].column_letter

    p_wb.close()
    o_wb.close()


def test_compat_append_and_iteration():
    p_wb = rpyxl.Workbook()
    o_wb = openpyxl.Workbook()

    p_ws = p_wb.active
    o_ws = o_wb.active

    rows = [
        ["ID", "Name", "Score", "Active"],
        [101, "Alice", 95.5, True],
        [102, "Bob", 88.0, False],
        [103, "Charlie", 91.25, True],
    ]

    for r in rows:
        p_ws.append(r)
        o_ws.append(r)

    assert p_ws.dimensions == o_ws.dimensions
    assert p_ws.min_row == o_ws.min_row
    assert p_ws.max_row == o_ws.max_row
    assert p_ws.min_column == o_ws.min_column
    assert p_ws.max_column == o_ws.max_column

    p_iter = list(p_ws.iter_rows(values_only=True))
    o_iter = list(o_ws.iter_rows(values_only=True))
    assert p_iter == o_iter

    p_wb.close()
    o_wb.close()


def test_compat_exception_behavior():
    p_wb = rpyxl.Workbook()
    o_wb = openpyxl.Workbook()

    # KeyError for missing worksheet
    with pytest.raises(KeyError):
        _ = p_wb["NonExistentSheet"]

    with pytest.raises(KeyError):
        _ = o_wb["NonExistentSheet"]

    p_wb.close()
    o_wb.close()
