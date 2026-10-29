"""Exact text comparisons across writers, readers, and streaming modes."""

import openpyxl
import pytest

import rpyxl

TEXT_VALUES = [
    "  padded  ",
    "   ",
    "\t padded\t",
    "line1\nline2",
    "line1\r\nline2",
    "line1\rline2",
    "héllo 日本語 😀 <&>\nline2",
    "literal _x000D_ text",
]


@pytest.mark.parametrize("writer", [openpyxl, rpyxl], ids=["openpyxl", "rpyxl"])
@pytest.mark.parametrize("write_only", [False, True], ids=["normal-write", "stream-write"])
@pytest.mark.parametrize("read_only", [False, True], ids=["normal-read", "stream-read"])
def test_text_fidelity_across_libraries(tmp_path, writer, write_only, read_only):
    path = tmp_path / "text.xlsx"
    workbook = writer.Workbook(write_only=write_only)
    try:
        sheet = workbook.create_sheet("Text") if write_only else workbook.active
        for value in TEXT_VALUES:
            sheet.append([value])
        workbook.save(str(path))
    finally:
        workbook.close()

    expected_values = TEXT_VALUES
    if writer is openpyxl:
        # Reference writers can emit literal CRs, which XML readers normalize.
        # Compare import behavior with the reference instead of assuming its output.
        reference = openpyxl.load_workbook(str(path))
        try:
            expected_values = [row[0] for row in reference.active.iter_rows(values_only=True)]
        finally:
            reference.close()
        for original, actual in zip(TEXT_VALUES, expected_values):
            if "\r" not in original:
                assert actual == original

    for reader in (openpyxl, rpyxl):
        loaded = reader.load_workbook(str(path), read_only=read_only)
        try:
            sheet = loaded.active
            values = [row[0] for row in sheet.iter_rows(values_only=True)]
            assert values == expected_values, f"{writer.__name__} -> {reader.__name__}"
            for row, expected in enumerate(expected_values, 1):
                assert sheet.cell(row=row, column=1).value == expected
        finally:
            loaded.close()
