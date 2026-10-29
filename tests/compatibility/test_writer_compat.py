"""Compatibility tests verifying round-trip fidelity between RPyXL, openpyxl, and calamine."""

import datetime
from pathlib import Path
import pytest
import openpyxl
from python_calamine import CalamineWorkbook

import rpyxl


class TestWriterCompatibility:
    def test_rpyxl_standard_write_openpyxl_read(self, tmp_path):
        file_path = tmp_path / "rpyxl_to_openpyxl.xlsx"

        # Create workbook with rpyxl
        wb = rpyxl.Workbook()
        ws = wb.active
        ws.title = "CompatTest"

        # Headers
        ws.append(["ID", "Name", "Score", "Active", "Date", "Formula"])

        # Rows
        d = datetime.date(2023, 11, 20)
        ws.append([1, "Alice", 95.5, True, d, "=SUM(A2:A2)"])
        ws.append([2, "Bob & Co <Ltd>", 88.0, False, d, "=A3*2"])
        ws.append([3, "日本語 / 中文 / 🌟", 99.9, True, d, "=A4+10"])

        wb.save(file_path)
        assert file_path.exists()

        # Read back with openpyxl
        ox_wb = openpyxl.load_workbook(file_path)
        assert ox_wb.sheetnames == ["CompatTest"]
        ox_ws = ox_wb["CompatTest"]

        assert ox_ws["A1"].value == "ID"
        assert ox_ws["B1"].value == "Name"
        assert ox_ws["C1"].value == "Score"
        assert ox_ws["D1"].value == "Active"

        assert ox_ws["A2"].value == 1
        assert ox_ws["B2"].value == "Alice"
        assert abs(ox_ws["C2"].value - 95.5) < 1e-4
        assert ox_ws["D2"].value is True

        assert ox_ws["B3"].value == "Bob & Co <Ltd>"
        assert ox_ws["B4"].value == "日本語 / 中文 / 🌟"

    def test_rpyxl_write_only_openpyxl_read(self, tmp_path):
        file_path = tmp_path / "rpyxl_wo_to_openpyxl.xlsx"

        wb = rpyxl.Workbook(write_only=True)
        ws = wb.create_sheet("Streaming")

        rows_count = 1000
        for i in range(1, rows_count + 1):
            ws.append([i, f"Product-{i}", i * 1.5, i % 2 == 0])

        wb.save(file_path)

        # Read back with openpyxl
        ox_wb = openpyxl.load_workbook(file_path, read_only=True)
        assert ox_wb.sheetnames == ["Streaming"]
        ox_ws = ox_wb["Streaming"]

        read_count = 0
        for r_idx, row in enumerate(ox_ws.iter_rows(values_only=True), start=1):
            assert row[0] == r_idx
            assert row[1] == f"Product-{r_idx}"
            assert abs(row[2] - r_idx * 1.5) < 1e-4
            assert row[3] == (r_idx % 2 == 0)
            read_count += 1

        assert read_count == rows_count

    def test_rpyxl_write_calamine_read(self, tmp_path):
        file_path = tmp_path / "rpyxl_to_calamine.xlsx"

        wb = rpyxl.Workbook()
        ws = wb.active
        ws["A1"] = 100
        ws["B1"] = "Calamine Engine"
        ws["C1"] = 2.71828
        ws.append([200, "Row 2", 1.41421])
        wb.save(file_path)

        cal = CalamineWorkbook.from_path(str(file_path))
        data = cal.get_sheet_by_name("Sheet").to_python()
        assert data[0][0] == 100.0
        assert data[0][1] == "Calamine Engine"
        assert abs(data[0][2] - 2.71828) < 1e-4
        assert data[1][0] == 200.0
        assert data[1][1] == "Row 2"

    def test_openpyxl_write_rpyxl_read(self, tmp_path):
        file_path = tmp_path / "openpyxl_to_rpyxl.xlsx"

        ox_wb = openpyxl.Workbook()
        ox_ws = ox_wb.active
        ox_ws.title = "SourceData"
        ox_ws["A1"] = "Created by openpyxl"
        ox_ws["B1"] = 12345
        ox_ws["C1"] = 45.67
        ox_ws["D1"] = True
        ox_wb.save(file_path)

        rpyxl_wb = rpyxl.load_workbook(file_path)
        assert rpyxl_wb.sheetnames == ["SourceData"]
        p_ws = rpyxl_wb["SourceData"]
        assert p_ws["A1"].value == "Created by openpyxl"
        assert p_ws["B1"].value == 12345
        assert abs(p_ws["C1"].value - 45.67) < 1e-4
        assert p_ws["D1"].value is True

    def test_openpyxl_write_only_to_rpyxl_streaming_read(self, tmp_path):
        file_path = tmp_path / "openpyxl_wo_to_rpyxl.xlsx"

        ox_wb = openpyxl.Workbook(write_only=True)
        ox_ws = ox_wb.create_sheet("BatchData")
        for i in range(1, 101):
            ox_ws.append([i, f"Batch-{i}"])
        ox_wb.save(file_path)

        rpyxl_wb = rpyxl.load_workbook(file_path, read_only=True)
        p_ws = rpyxl_wb["BatchData"]
        rows = list(p_ws.iter_rows(values_only=True))
        assert len(rows) == 100
        assert rows[0] == (1, "Batch-1")
        assert rows[99] == (100, "Batch-100")
