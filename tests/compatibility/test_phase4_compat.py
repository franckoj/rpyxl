"""Comprehensive Phase 4 openpyxl behavioral compatibility test suite."""

import datetime
import pytest
import rpyxl
from rpyxl.cell import (
    TYPE_STRING,
    TYPE_NUMERIC,
    TYPE_BOOL,
    TYPE_NULL,
    TYPE_FORMULA,
    TYPE_DATE,
)
from rpyxl.utils import (
    get_column_letter,
    column_index_from_string,
    coordinate_to_tuple,
    range_boundaries,
    quote_sheetname,
)
from rpyxl.utils.exceptions import (
    OpenpyxlException,
    InvalidFileException,
    SheetTitleException,
    CellCoordinatesException,
    NamedRangeException,
)
from rpyxl.styles import Font, PatternFill, Alignment, Border, Side, Color

import openpyxl
from openpyxl.utils import (
    range_boundaries as o_range_boundaries,
    quote_sheetname as o_quote_sheetname,
)


class TestWorkbookCompatibility:
    """Test Workbook methods and behaviors against openpyxl."""

    def test_workbook_initialization_and_properties(self):
        p_wb = rpyxl.Workbook()
        o_wb = openpyxl.Workbook()

        assert p_wb.sheetnames == o_wb.sheetnames
        assert p_wb.active.title == o_wb.active.title
        assert len(p_wb.worksheets) == len(o_wb.worksheets)
        assert p_wb.read_only == o_wb.read_only
        assert p_wb.write_only == o_wb.write_only
        assert p_wb.data_only == o_wb.data_only
        assert len(p_wb.sheetnames) == len(o_wb.sheetnames)

        p_wb.close()
        o_wb.close()

    def test_active_sheet_assignment_index_and_object(self):
        p_wb = rpyxl.Workbook()
        o_wb = openpyxl.Workbook()

        p_ws1 = p_wb.create_sheet("Sheet2")
        o_ws1 = o_wb.create_sheet("Sheet2")

        # Set active by index
        p_wb.active = 1
        o_wb.active = 1
        assert p_wb.active.title == o_wb.active.title

        # Set active by worksheet object
        p_wb.active = p_wb["Sheet"]
        o_wb.active = o_wb["Sheet"]
        assert p_wb.active.title == o_wb.active.title

        p_wb.close()
        o_wb.close()

    def test_copy_worksheet(self):
        p_wb = rpyxl.Workbook()
        o_wb = openpyxl.Workbook()

        p_wb.active["A1"] = "Data123"
        o_wb.active["A1"] = "Data123"

        p_copy = p_wb.copy_worksheet(p_wb.active)
        o_copy = o_wb.copy_worksheet(o_wb.active)

        assert p_copy.title == o_copy.title
        assert p_copy["A1"].value == o_copy["A1"].value
        assert p_wb.sheetnames == o_wb.sheetnames

        p_wb.close()
        o_wb.close()

    def test_workbook_index(self):
        p_wb = rpyxl.Workbook()
        o_wb = openpyxl.Workbook()

        p_ws2 = p_wb.create_sheet("Second")
        o_ws2 = o_wb.create_sheet("Second")

        assert p_wb.index(p_ws2) == o_wb.index(o_ws2)

        other_wb = rpyxl.Workbook()
        other_ws = other_wb.active
        with pytest.raises(ValueError):
            p_wb.index(other_ws)

        p_wb.close()
        o_wb.close()
        other_wb.close()


class TestWorksheetCompatibility:
    """Test Worksheet slicing, indexing, appending, and iteration against openpyxl."""

    def test_cell_arguments_and_validation(self):
        p_wb = rpyxl.Workbook()
        o_wb = openpyxl.Workbook()

        p_ws = p_wb.active
        o_ws = o_wb.active

        # Positional arguments: (row, column)
        p_c1 = p_ws.cell(1, 2)
        o_c1 = o_ws.cell(1, 2)
        assert p_c1.coordinate == o_c1.coordinate

        # Keyword arguments: (row=..., column=...)
        p_c2 = p_ws.cell(row=2, column=3)
        o_c2 = o_ws.cell(row=2, column=3)
        assert p_c2.coordinate == o_c2.coordinate

        # Positional with value: (row, column, value)
        p_c3 = p_ws.cell(3, 4, "pos_val")
        o_c3 = o_ws.cell(3, 4, "pos_val")
        assert p_c3.coordinate == o_c3.coordinate
        assert p_c3.value == o_c3.value

        # Keyword with value: (row=..., column=..., value=...)
        p_c4 = p_ws.cell(row=4, column=5, value="kw_val")
        o_c4 = o_ws.cell(row=4, column=5, value="kw_val")
        assert p_c4.coordinate == o_c4.coordinate
        assert p_c4.value == o_c4.value

        # Zero and negative indices raise ValueError
        with pytest.raises(ValueError):
            p_ws.cell(0, 1)
        with pytest.raises(ValueError):
            o_ws.cell(0, 1)

        with pytest.raises(ValueError):
            p_ws.cell(1, 0)
        with pytest.raises(ValueError):
            o_ws.cell(1, 0)

        p_wb.close()
        o_wb.close()

    def test_append_varieties_including_dict(self):
        p_wb = rpyxl.Workbook()
        o_wb = openpyxl.Workbook()

        p_ws = p_wb.active
        o_ws = o_wb.active

        # List
        p_ws.append(["A", "B", "C"])
        o_ws.append(["A", "B", "C"])

        # Tuple
        p_ws.append((10, 20, 30))
        o_ws.append((10, 20, 30))

        # Generator
        p_ws.append(x * 2 for x in [1, 2, 3])
        o_ws.append(x * 2 for x in [1, 2, 3])

        # Dict with column letters
        p_ws.append({"A": "DictColA", "C": "DictColC"})
        o_ws.append({"A": "DictColA", "C": "DictColC"})

        # Dict with integer columns
        p_ws.append({1: "DictCol1", 4: "DictCol4"})
        o_ws.append({1: "DictCol1", 4: "DictCol4"})

        assert list(p_ws.iter_rows(values_only=True)) == list(o_ws.iter_rows(values_only=True))

        p_wb.close()
        o_wb.close()

    def test_slicing_and_indexing_types(self):
        p_wb = rpyxl.Workbook()
        o_wb = openpyxl.Workbook()

        p_ws = p_wb.active
        o_ws = o_wb.active

        for r in range(1, 4):
            for c in range(1, 4):
                val = f"R{r}C{c}"
                p_ws.cell(r, c, value=val)
                o_ws.cell(r, c, value=val)

        # Single cell access
        assert p_ws["A1"].value == o_ws["A1"].value
        assert p_ws["B2"].value == o_ws["B2"].value

        # Integer row indexing: ws[1]
        p_row1 = [c.value for c in p_ws[1]]
        o_row1 = [c.value for c in o_ws[1]]
        assert p_row1 == o_row1

        # Pure column letter: ws["A"]
        p_colA = [c.value for c in p_ws["A"]]
        o_colA = [c.value for c in o_ws["A"]]
        assert p_colA == o_colA

        # Column range: ws["A:B"]
        p_colAB = [[c.value for c in col] for col in p_ws["A:B"]]
        o_colAB = [[c.value for c in col] for col in o_ws["A:B"]]
        assert p_colAB == o_colAB

        # Row range string: ws["1:2"]
        p_row12_str = [[c.value for c in row] for row in p_ws["1:2"]]
        o_row12_str = [[c.value for c in row] for row in o_ws["1:2"]]
        assert p_row12_str == o_row12_str

        # Row range slice: ws[1:2]
        p_row12_slice = [[c.value for c in row] for row in p_ws[1:2]]
        o_row12_slice = [[c.value for c in row] for row in o_ws[1:2]]
        assert p_row12_slice == o_row12_slice

        # Rectangular range string: ws["A1:B2"]
        p_box_str = [[c.value for c in row] for row in p_ws["A1:B2"]]
        o_box_str = [[c.value for c in row] for row in o_ws["A1:B2"]]
        assert p_box_str == o_box_str

        # Rectangular range slice: ws["A1":"B2"]
        p_box_slice = [[c.value for c in row] for row in p_ws["A1":"B2"]]
        o_box_slice = [[c.value for c in row] for row in o_ws["A1":"B2"]]
        assert p_box_slice == o_box_slice

        p_wb.close()
        o_wb.close()

    def test_iter_cols_positional_and_keyword(self):
        p_wb = rpyxl.Workbook()
        o_wb = openpyxl.Workbook()

        p_ws = p_wb.active
        o_ws = o_wb.active

        for r in range(1, 5):
            for c in range(1, 5):
                val = r * 10 + c
                p_ws.cell(r, c, value=val)
                o_ws.cell(r, c, value=val)

        # iter_cols positional: (min_col, max_col, min_row, max_row, values_only)
        p_cols_pos = list(p_ws.iter_cols(2, 3, 2, 3, True))
        o_cols_pos = list(o_ws.iter_cols(2, 3, 2, 3, True))
        assert p_cols_pos == o_cols_pos

        # iter_cols keyword
        p_cols_kw = list(p_ws.iter_cols(min_col=1, max_col=2, min_row=1, max_row=2, values_only=True))
        o_cols_kw = list(o_ws.iter_cols(min_col=1, max_col=2, min_row=1, max_row=2, values_only=True))
        assert p_cols_kw == o_cols_kw

        p_wb.close()
        o_wb.close()

    def test_empty_sheet_bounds_and_properties(self):
        p_wb = rpyxl.Workbook()
        o_wb = openpyxl.Workbook()

        p_ws = p_wb.active
        o_ws = o_wb.active

        assert p_ws.min_row == o_ws.min_row
        assert p_ws.max_row == o_ws.max_row
        assert p_ws.min_column == o_ws.min_column
        assert p_ws.max_column == o_ws.max_column

        assert list(p_ws.rows) == list(o_ws.rows)
        assert list(p_ws.columns) == list(o_ws.columns)
        assert list(p_ws.values) == list(o_ws.values)

        p_wb.close()
        o_wb.close()

    def test_parent_references(self):
        p_wb = rpyxl.Workbook()
        p_ws = p_wb.active
        p_cell = p_ws["A1"]

        assert p_ws.parent.sheetnames == p_wb.sheetnames
        assert p_cell.parent.title == p_ws.title

        p_wb.close()


class TestCellCompatibility:
    """Test Cell attributes, formats, data types, and setters."""

    def test_cell_data_types_and_date_check(self):
        p_wb = rpyxl.Workbook()
        o_wb = openpyxl.Workbook()

        p_ws = p_wb.active
        o_ws = o_wb.active

        test_cases = [
            ("A1", "hello", "s", False),
            ("A2", 42, "n", False),
            ("A3", 3.14, "n", False),
            ("A4", True, "b", False),
            ("A5", datetime.date(2023, 1, 1), "d", True),
            ("A6", datetime.datetime(2023, 1, 1, 12, 0, 0), "d", True),
            ("A7", "=SUM(A2:A3)", "f", False),
        ]

        for coord, val, expected_type, expected_is_date in test_cases:
            p_ws[coord] = val
            o_ws[coord] = val

            p_cell = p_ws[coord]
            o_cell = o_ws[coord]

            assert p_cell.data_type == expected_type
            assert p_cell.is_date == expected_is_date
            assert o_cell.is_date == expected_is_date
            assert p_cell.value == o_cell.value

        p_wb.close()
        o_wb.close()

    def test_cell_styling_properties(self):
        wb = rpyxl.Workbook()
        ws = wb.active
        cell = ws["A1"]

        assert cell.number_format == "General"
        cell.number_format = "#,##0.00"
        assert cell.number_format == "#,##0.00"

        assert cell.comment is None
        cell.comment = "Review this figure"
        assert cell.comment == "Review this figure"

        assert cell.hyperlink is None
        cell.hyperlink = "https://example.com"
        assert cell.hyperlink == "https://example.com"

        wb.close()

    def test_cell_type_constants(self):
        assert TYPE_STRING == "s"
        assert TYPE_NUMERIC == "n"
        assert TYPE_BOOL == "b"
        assert TYPE_NULL == "n"
        assert TYPE_FORMULA == "f"
        assert TYPE_DATE == "d"


class TestUtilsAndStylesFoundation:
    """Test utils and styles modules."""

    def test_coordinate_utils(self):
        assert get_column_letter(1) == "A"
        assert get_column_letter(28) == "AB"
        assert column_index_from_string("A") == 1
        assert column_index_from_string("AB") == 28
        assert coordinate_to_tuple("B5") == (5, 2)

        assert range_boundaries("A1:C5") == o_range_boundaries("A1:C5")
        assert range_boundaries("B3") == o_range_boundaries("B3")

        assert quote_sheetname("Sheet 1") == o_quote_sheetname("Sheet 1")
        assert quote_sheetname("Simple") == o_quote_sheetname("Simple")

    def test_exceptions_hierarchy(self):
        assert issubclass(InvalidFileException, OpenpyxlException)
        assert issubclass(SheetTitleException, (OpenpyxlException, ValueError))
        assert issubclass(CellCoordinatesException, (OpenpyxlException, ValueError))
        assert issubclass(NamedRangeException, OpenpyxlException)

    def test_styles_foundation(self):
        font = Font(name="Arial", size=12, bold=True, italic=True)
        assert font.name == "Arial"
        assert font.size == 12
        assert font.bold is True

        fill = PatternFill(fill_type="solid", start_color="FFFF0000")
        assert fill.fill_type == "solid"
        assert fill.start_color == "FFFF0000"

        align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        assert align.horizontal == "center"
        assert align.wrap_text is True

        side = Side(style="thin", color="000000")
        border = Border(left=side, right=side, top=side, bottom=side)
        assert border.left.style == "thin"

        color = Color(rgb="FF0000")
        assert color.rgb == "FF0000"
